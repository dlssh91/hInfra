"""서버(OS) 기준 로더 + main.run 통합 테스트 (작업③ — 합성 픽스처 기반).

주의: 합성 xlsx/XML 회귀 테스트이며, 실수집 데이터 대상 e2e는 샘플 미확보로
보류 상태다(샘플 확보 후 별도 진행). results/ 실데이터는 사용하지 않는다.
"""
import json

import openpyxl
import pytest

from judge_tool.criteria_loader import load_criteria
from judge_tool.errors import ReportError
from judge_tool.main import run
from judge_tool.profile import SERVER

# 변형별 (평가대상 'o', 판단기준, 판단방법) 컬럼 — profile.SERVER와 동일 배치.
_COLS = {
    "aix": (12, 17, 18),
    "hpux": (13, 19, 20),
    "linux": (14, 21, 22),
    "solaris": (15, 23, 24),
    "win": (16, 25, 26),
}


def _server_xlsx(path, rows):
    """합성 '서버' 시트. rows: [(item_id, {variant: applicable_bool})]."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "서버"
    ws.cell(4, 2, "평가항목ID"); ws.cell(4, 7, "평가항목"); ws.cell(4, 8, "위험도")
    for vname, (app, std, mth) in _COLS.items():
        ws.cell(4, app, f"평가대상({vname})")
        ws.cell(4, std, f"판단기준({vname})")
        ws.cell(4, mth, f"판단방법({vname})")
    r = 5
    for item_id, applicables in rows:
        ws.cell(r, 2, item_id); ws.cell(r, 7, f"{item_id}항목"); ws.cell(r, 8, 5.0)
        for vname, (app, std, mth) in _COLS.items():
            if applicables.get(vname):
                ws.cell(r, app, "o")
            ws.cell(r, std, "* 양호 - ...\n* 취약 - ...")
            ws.cell(r, mth, "확인방법")
        r += 1
    wb.save(path)


def test_server_criteria_all_variants(tmp_path):
    p = str(tmp_path / "server.xlsx")
    _server_xlsx(p, [
        ("SRV-001", {v: True for v in _COLS}),       # 전 OS 적용
        ("SRV-002", {"linux": True}),                 # linux만 적용
    ])
    crit = load_criteria(p, SERVER, profile_key="server")

    for vname in _COLS:
        c = crit[("SRV-001", vname)]
        assert c.applicable is True and c.is_judgeable is True
        assert "양호" in c.standard
        # Phase 1: server.yaml에 SRV-001 det_common 등재 → label A, det_common
        assert c.label == "A"
        assert c.judgment_method == "det_common"

    assert crit[("SRV-002", "linux")].is_judgeable is True
    for vname in ("aix", "hpux", "solaris", "win"):
        c = crit[("SRV-002", vname)]
        assert c.applicable is False and c.is_judgeable is False


# ── main.run 통합(가벼운 스텁) ───────────────────────────────────────────────

class StubVuln:
    def chat(self, system, user):
        return ('{"verdict":"취약","confidence":0.8,'
                '"rationale":"테스트","cited_evidence":["x"]}')


_SERVER_XML = """<?xml version="1.0" encoding="UTF-8"?>
<?xml-stylesheet type="text/xsl" href="isac.xsl"?>
<script>
<asset>
<hostname>testhost</hostname>
<os>Linux</os>
<uname>Linux testhost 5.14.0</uname>
<whoami>root</whoami>
<version>2026.1</version>
</asset>
<results>
<dump>
<items><id>SRV-001</id></items>
<output><![CDATA[PermitRootLogin no]]></output>
</dump>
<dump>
<items><id>SRV-002</id></items>
<output><![CDATA[   ]]></output>
</dump>
<dump>
<items><id>Internet</id></items>
<output><![CDATA[ping ok]]></output>
</dump>
</results>
</script>
"""


def test_server_run_integration_detects_variant(tmp_path):
    """합성 서버 xlsx + 합성 서버 XML(파일명에 OS 마커 없음) + 스텁 LLM으로
    run() 배선을 검증: 내용 기반(detect_variant) 변형 자동식별 → SRV 항목
    판정 → json/xlsx 산출 + judgment_method 노출.

    주의: 합성 픽스처 기반 회귀 테스트다. 실수집 데이터 대상 e2e는
    샘플 데이터 미확보로 보류(후속 과제)."""
    criteria = str(tmp_path / "server.xlsx")
    _server_xlsx(criteria, [
        ("SRV-001", {v: True for v in _COLS}),
        ("SRV-002", {v: True for v in _COLS}),
    ])
    # 실제 수집 파일명 형식({hostname}-s-{date}.xml) — OS 변형 마커 없음.
    report = str(tmp_path / "testhost-s-20260612.xml")
    with open(report, "w", encoding="utf-8") as fh:
        fh.write(_SERVER_XML)
    jout = str(tmp_path / "r.json"); xout = str(tmp_path / "r.xlsx")
    cov = run(report, criteria, "server", StubVuln(), jout, xout, "stub")

    data = json.load(open(jout, encoding="utf-8"))
    # 파일명 폴백: <asset><os>Linux → linux 변형 자동식별
    assert data["metadata"]["variant"] == "linux"
    assert data["metadata"]["profile"] == "server"
    by_id = {j["item_id"]: j for j in data["judgments"]}
    # SRV-001: Phase 1 server.yaml → det_common. raw 증거("PermitRootLogin no")로
    #   check_SRV_001 결정론 판정 → 결과는 결정론에 따름(양호/취약).
    #   판정 성공 시 needs_review=True(결정론 정당성=사람 검토).
    #   판정 실패(handled=False) 시 §18.3 라벨A → LLM 폴백.
    assert by_id["SRV-001"]["needs_review"] is True
    assert by_id["SRV-001"]["judgment_method"] == "det_common"
    assert by_id["SRV-001"]["script_status"] is None    # status_available=False
    assert by_id["SRV-001"]["agreement"] == "N/A"
    # SRV-002: 공백 output → 증거 없음 → 판단보류 강제(empty_means_good 없음)
    assert by_id["SRV-002"]["verdict"] == "판단보류"
    # 비-SRV 덤프(Internet)는 기준에 없어 자연 스킵
    assert "INTERNET" not in by_id and "Internet" not in by_id
    assert cov["judged"] == 2 and cov["expected"] == 2
    # xlsx 산출 확인
    wb = openpyxl.load_workbook(xout)
    assert "판정결과" in wb.sheetnames


def test_server_run_unknown_os_guides_variant_flag(tmp_path):
    """미지 OS(detect_variant=None) → --variant 안내 ReportError
    (서버는 파일명 마커가 없으므로 내용 기반 식별 실패 메시지)."""
    criteria = str(tmp_path / "server.xlsx")
    _server_xlsx(criteria, [("SRV-001", {v: True for v in _COLS})])
    report = str(tmp_path / "testhost-s-1.xml")
    with open(report, "w", encoding="utf-8") as fh:
        fh.write(_SERVER_XML.replace("<os>Linux</os>", "<os>FreeBSD</os>"))
    with pytest.raises(ReportError, match="--variant"):
        run(report, criteria, "server", StubVuln(),
            str(tmp_path / "j.json"), str(tmp_path / "r.xlsx"), "stub")


def test_server_run_variant_override_skips_detection(tmp_path):
    """--variant 오버라이드는 내용 기반 식별보다 우선한다."""
    criteria = str(tmp_path / "server.xlsx")
    _server_xlsx(criteria, [("SRV-001", {v: True for v in _COLS})])
    report = str(tmp_path / "testhost-s-2.xml")
    with open(report, "w", encoding="utf-8") as fh:
        fh.write(_SERVER_XML)   # <os>Linux</os> 이지만 aix로 강제
    jout = str(tmp_path / "j.json")
    run(report, criteria, "server", StubVuln(), jout,
        str(tmp_path / "r.xlsx"), "stub", variant_override="aix")
    data = json.load(open(jout, encoding="utf-8"))
    assert data["metadata"]["variant"] == "aix"
