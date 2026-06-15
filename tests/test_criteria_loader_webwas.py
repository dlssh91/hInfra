"""웹서버-WAS(webwas 프로파일) 기준 로더 단위 테스트.

합성 xlsx 픽스처 기반 — 실수집 데이터 없이 실행한다.
핵심 검증: (1) OS 변형 전용 판단컬럼(col23~32, 기준 먼저),
           (2) 웹서버 변형 공통 판단컬럼(col37/38) 공유,
           (3) 변형별 적용 차이(OS 5종 / 웹서버 6종).
"""
import openpyxl
import pytest

from judge_tool.criteria_loader import load_criteria
from judge_tool.profile import WEBWAS

# 컬럼 배치 (1-indexed, openpyxl ws.cell 기준)
_ID, _NAME, _RISK = 2, 7, 8
_APP = {"aix": 12, "hpux": 13, "linux": 14, "solaris": 15, "win": 16,
        "webservice": 17, "apache": 18, "webtob": 19, "iis": 20,
        "tomcat": 21, "jeus": 22}
# OS 변형 전용 판단컬럼 (기준, 방법)
_OS_STD_MTH = {"aix": (23, 24), "hpux": (25, 26), "linux": (27, 28),
               "solaris": (29, 30), "win": (31, 32)}
_COMMON_STD, _COMMON_MTH = 37, 38


def _webwas_xlsx(tmp_path, rows):
    """합성 '웹서버-WAS' 시트.

    rows: [(item_id, {variant: bool 적용}, os_std, os_mth, common_std, common_mth)]
    os_std/os_mth는 OS 변형 전용 컬럼에, common_*는 공통 col37/38에 채운다.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "웹서버-WAS"
    ws.cell(4, _ID, "평가항목ID")
    ws.cell(4, _NAME, "평가항목")
    ws.cell(4, _RISK, "위험도")
    for v, col in _APP.items():
        ws.cell(4, col, f"평가대상({v})")
    for v, (sc, mc) in _OS_STD_MTH.items():
        ws.cell(4, sc, f"판단기준({v})")
        ws.cell(4, mc, f"판단방법({v})")
    ws.cell(4, _COMMON_STD, "판단기준")
    ws.cell(4, _COMMON_MTH, "판단방법")
    r = 5
    for item_id, applies, os_std, os_mth, common_std, common_mth in rows:
        ws.cell(r, _ID, item_id)
        ws.cell(r, _NAME, f"{item_id}항목")
        ws.cell(r, _RISK, 4.0)
        for v, on in applies.items():
            if on:
                ws.cell(r, _APP[v], "o")
        # OS 전용 컬럼 채움
        for v, (sc, mc) in _OS_STD_MTH.items():
            if os_std:
                ws.cell(r, sc, os_std)
            if os_mth:
                ws.cell(r, mc, os_mth)
        if common_std:
            ws.cell(r, _COMMON_STD, common_std)
        if common_mth:
            ws.cell(r, _COMMON_MTH, common_mth)
        r += 1
    p = str(tmp_path / "webwas.xlsx")
    wb.save(p)
    return p


# ── OS 변형 전용 컬럼 ─────────────────────────────────────────────────────────

def test_os_variant_uses_dedicated_columns(tmp_path):
    """OS 항목(WST-001): OS 변형이 전용 판단컬럼(기준 먼저, 방법 나중) 사용."""
    p = _webwas_xlsx(tmp_path, [
        ("WST-001", {"aix": True, "linux": True}, "OS_기준_텍스트", "OS_방법_텍스트",
         None, None),
    ])
    crit = load_criteria(p, WEBWAS, profile_key="webwas")
    for v in ("aix", "linux"):
        c = crit[("WST-001", v)]
        assert c.applicable is True
        assert c.is_judgeable is True
        assert c.standard == "OS_기준_텍스트", f"{v}: standard 오매핑"
        assert c.method == "OS_방법_텍스트", f"{v}: method 오매핑"
        assert c.judgment_method == "llm"


def test_os_variant_not_applicable_when_no_o(tmp_path):
    """WST-001에서 solaris 'o' 없음 → not applicable."""
    p = _webwas_xlsx(tmp_path, [
        ("WST-001", {"aix": True}, "OS_기준", "OS_방법", None, None),
    ])
    crit = load_criteria(p, WEBWAS, profile_key="webwas")
    assert crit[("WST-001", "solaris")].applicable is False


# ── 웹서버 변형 공통 컬럼 ─────────────────────────────────────────────────────

def test_web_variant_uses_common_columns(tmp_path):
    """웹 특화(WST-031): 웹서버 변형이 공통 col37/38 사용."""
    p = _webwas_xlsx(tmp_path, [
        ("WST-031",
         {"linux": True, "apache": True, "tomcat": True, "iis": True},
         "OS_기준", "OS_방법", "공통_기준_텍스트", "공통_방법_텍스트"),
    ])
    crit = load_criteria(p, WEBWAS, profile_key="webwas")
    for v in ("apache", "tomcat", "iis"):
        c = crit[("WST-031", v)]
        assert c.applicable is True
        assert c.is_judgeable is True
        assert c.standard == "공통_기준_텍스트", f"{v}: 공통 standard 오매핑"
        assert c.method == "공통_방법_텍스트", f"{v}: 공통 method 오매핑"


def test_web_variant_jeus_not_applicable(tmp_path):
    """WST-031에서 jeus 'o' 없음 → not applicable (웹서버 변형별 적용 차이)."""
    p = _webwas_xlsx(tmp_path, [
        ("WST-031", {"linux": True, "apache": True}, "OS_기준", "OS_방법",
         "공통_기준", "공통_방법"),
    ])
    crit = load_criteria(p, WEBWAS, profile_key="webwas")
    assert crit[("WST-031", "jeus")].applicable is False


def test_web_variant_no_common_standard_not_judgeable(tmp_path):
    """웹서버 변형 'o'지만 공통 판단기준 빈칸 → is_judgeable=False."""
    p = _webwas_xlsx(tmp_path, [
        ("WST-031", {"apache": True}, "OS_기준", "OS_방법", "", ""),
    ])
    crit = load_criteria(p, WEBWAS, profile_key="webwas")
    c = crit[("WST-031", "apache")]
    assert c.applicable is True
    assert c.is_judgeable is False


# ── 프로파일 등록 ─────────────────────────────────────────────────────────────

def test_webwas_profile_registered():
    """get_profile('webwas'): parser=webwas_xml, 11변형, 컬럼 비대칭 매핑."""
    from judge_tool.profile import get_profile
    prof = get_profile("webwas")
    assert prof.parser == "webwas_xml"
    assert set(prof.variants.keys()) == {
        "aix", "hpux", "linux", "solaris", "win",
        "webservice", "apache", "webtob", "iis", "tomcat", "jeus"}
    # OS 변형: 전용 컬럼 (기준 먼저)
    assert prof.variants["aix"].standard_col == 23
    assert prof.variants["aix"].method_col == 24
    assert prof.variants["win"].standard_col == 31
    assert prof.variants["win"].method_col == 32
    # 웹서버 변형: 공통 컬럼 공유
    for v in ("webservice", "apache", "webtob", "iis", "tomcat", "jeus"):
        assert prof.variants[v].standard_col == 37, f"{v} standard_col"
        assert prof.variants[v].method_col == 38, f"{v} method_col"
    # applicability_col 갈림
    assert prof.variants["apache"].applicability_col == 18
    assert prof.variants["jeus"].applicability_col == 22
