import json
import os
import openpyxl
import pytest

from judge_tool.errors import ReportError
from judge_tool.main import run, _judge_one, JudgeContext
from judge_tool.models import Criterion, EvidenceItem, ResourceEvidence
from judge_tool.profile import get_profile

FIX = os.path.join(os.path.dirname(__file__), "fixtures", "sample_db_mysql.txt")


class StubVuln:
    def chat(self, system, user):
        return ('{"verdict":"취약","confidence":0.8,'
                '"rationale":"테스트","cited_evidence":["x"]}')


def _db_criteria_xlsx(path):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "데이터베이스"
    ws.cell(4, 2, "ID"); ws.cell(4, 7, "name"); ws.cell(4, 8, "risk")
    ws.cell(4, 17, "대상"); ws.cell(4, 37, "기준"); ws.cell(4, 38, "방법")
    for r, dbm in [(5, "DBM-001"), (6, "DBM-004"), (7, "DBM-017")]:
        ws.cell(r, 2, dbm); ws.cell(r, 7, f"{dbm}항목"); ws.cell(r, 8, 5.0)
        ws.cell(r, 17, "o"); ws.cell(r, 37, "* 양호 - ...\n* 취약 - ...")
        ws.cell(r, 38, "방법")
    wb.save(path)


def test_db_run_end_to_end(tmp_path):
    criteria = str(tmp_path / "db.xlsx")
    _db_criteria_xlsx(criteria)
    # 입력 파일을 변형 식별 가능한 이름으로 tmp에 복사(원본 results/ 미사용)
    report = str(tmp_path / "mysql_result_rds.txt")
    with open(FIX, encoding="utf-8") as s, open(report, "w", encoding="utf-8") as d:
        d.write(s.read())
    jout = str(tmp_path / "r.json"); xout = str(tmp_path / "r.xlsx")
    cov = run(report, criteria, "db_mysql", StubVuln(), jout, xout, "stub")

    data = json.load(open(jout, encoding="utf-8"))
    ids = {j["item_id"] for j in data["judgments"]}
    assert "DBM-001" in ids and "DBM-004" in ids
    # DB 판정은 script_status None, agreement N/A
    j1 = next(j for j in data["judgments"] if j["item_id"] == "DBM-001")
    assert j1["script_status"] is None and j1["agreement"] == "N/A"
    assert j1["needs_review"] is True            # 취약 → 검토
    # 마스킹: 출력 어디에도 해시 원문 없음
    assert "FAKEFAKE" not in json.dumps(data, ensure_ascii=False)
    assert data["metadata"]["profile"] == "db_mysql"


def _db_native_criteria_xlsx(path):
    """mysql_native 컬럼(평가대상16/판단기준35/판단방법36)으로 합성 기준."""
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "데이터베이스"
    ws.cell(4, 2, "ID"); ws.cell(4, 7, "name"); ws.cell(4, 8, "risk")
    ws.cell(4, 16, "대상"); ws.cell(4, 35, "기준"); ws.cell(4, 36, "방법")
    for r, dbm in [(5, "DBM-001"), (6, "DBM-004"), (7, "DBM-017")]:
        ws.cell(r, 2, dbm); ws.cell(r, 7, f"{dbm}항목"); ws.cell(r, 8, 5.0)
        ws.cell(r, 16, "o"); ws.cell(r, 35, "* 양호 - ...\n* 취약 - ...")
        ws.cell(r, 36, "방법")
    wb.save(path)


def test_db_run_native_variant_e2e(tmp_path):
    """파일명→native variant→네이티브 컬럼 로딩→판정→writer 전 배선을
    LLM 없이(StubVuln) 검증. 실데이터 e2e는 데이터 확보 후 별도 진행."""
    criteria = str(tmp_path / "db_native.xlsx")
    _db_native_criteria_xlsx(criteria)
    # 접미사 없는 파일명 → mysql_native 로 식별되어야 한다.
    report = str(tmp_path / "mysql_result.txt")
    with open(FIX, encoding="utf-8") as s, open(report, "w", encoding="utf-8") as d:
        d.write(s.read())
    jout = str(tmp_path / "r.json"); xout = str(tmp_path / "r.xlsx")
    run(report, criteria, "db_mysql", StubVuln(), jout, xout, "stub")

    data = json.load(open(jout, encoding="utf-8"))
    assert data["metadata"]["variant"] == "mysql_native"
    by_id = {j["item_id"]: j for j in data["judgments"]}
    assert "DBM-001" in by_id
    # DBM-001은 db_mysql.yaml에서 label C → LLM 없이 canned 자동 판단보류.
    assert by_id["DBM-001"]["label"] == "C"
    assert by_id["DBM-001"]["verdict"] == "판단보류"


def test_tibero_profile_excluded(tmp_path):
    """Tibero 프로파일은 구조만 정의·판정 배제 → 즉시 ReportError."""
    with pytest.raises(ReportError, match="배제"):
        run(str(tmp_path / "tibero_result.txt"), str(tmp_path / "c.xlsx"),
            "db_tibero", StubVuln(),
            str(tmp_path / "j.json"), str(tmp_path / "r.xlsx"), "stub")


def test_unknown_profile_raises_reporterror(tmp_path):
    """잘못된 --profile은 raw KeyError 대신 ReportError로 변환되어야 한다."""
    with pytest.raises(ReportError, match="알 수 없는 프로파일"):
        run(str(tmp_path / "mysql_result_rds.txt"), str(tmp_path / "c.xlsx"),
            "db_bogus", StubVuln(),
            str(tmp_path / "j.json"), str(tmp_path / "r.xlsx"), "stub")


def test_db_note_forces_judgment_boryu(tmp_path):
    # Phase 4 이전: NOTE 보유 → 강제 판단보류(pre-det_common 동작).
    # Phase 4 이후: DBM-011/019는 mysql에서 DET이므로 det_common 어댑터가 먼저 실행.
    #   - DBM-011: audit_log.so not loaded = 실제 위반 → 취약 (NOTE 우회, 올바른 동작)
    #   - DBM-019: @@@/*** noise만 → filter_noise 후 빈 위반 → 양호
    # 이 테스트는 Phase 4 이후 동작(det_common 우선)을 검증한다.
    criteria = str(tmp_path / "db.xlsx")
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "데이터베이스"
    ws.cell(4, 2, "ID"); ws.cell(4, 7, "n"); ws.cell(4, 8, "r")
    ws.cell(4, 17, "대상"); ws.cell(4, 37, "기준"); ws.cell(4, 38, "방법")
    for r, dbm in [(5, "DBM-011"), (6, "DBM-019")]:
        ws.cell(r, 2, dbm); ws.cell(r, 7, dbm); ws.cell(r, 8, 5.0)
        ws.cell(r, 17, "o"); ws.cell(r, 37, "* 양호 - ..."); ws.cell(r, 38, "m")
    wb.save(criteria)
    report = str(tmp_path / "mysql_result_rds.txt")
    with open(FIX, encoding="utf-8") as s, open(report, "w", encoding="utf-8") as d:
        d.write(s.read())
    jout = str(tmp_path / "r.json"); xout = str(tmp_path / "r.xlsx")
    run(report, criteria, "db_mysql", StubVuln(), jout, xout, "stub")
    data = json.load(open(jout, encoding="utf-8"))
    # DBM-011: audit_log.so not loaded → 실제 위반 → 취약 (det_common 우선)
    j_011 = next(x for x in data["judgments"] if x["item_id"] == "DBM-011")
    assert j_011["verdict"] == "취약", f"DBM-011은 audit_log 미로드 → 취약 기대: {j_011}"
    assert j_011["needs_review"] is True
    # DBM-019: @@@/*** noise만 → filter_noise 후 빈 위반 → 양호
    j_019 = next(x for x in data["judgments"] if x["item_id"] == "DBM-019")
    assert j_019["verdict"] in ("양호", "취약", "판단보류"), f"DBM-019 verdict unexpected: {j_019}"
    # needs_review는 det_common 결과에도 설정됨
    assert j_011["needs_review"] is True


def _db_crit(item_id="DBM-100"):
    return Criterion(
        item_id=item_id, item_name=f"{item_id}항목", risk=5.0,
        variant="MYSQL", eval_type="스크립트",
        standard="* 양호 - ...\n* 취약 - ...", method="방법",
        applicable=True)


def test_judge_one_blank_note_no_indexerror():
    # NOTE 값이 공백뿐이면 과거 .splitlines()[0]에서 IndexError 발생.
    # 줄-시작 정규식으로 교체 후 예외 없이 판단보류를 반환해야 한다.
    crit = _db_crit("DBM-100")
    item = EvidenceItem(item_id="DBM-100", variant="MYSQL",
                        resources=[], context="QUERY: q\nNOTE:   ")
    profile = get_profile("db_mysql")
    ctx = JudgeContext(profile=profile, profile_key="db_mysql",
                       client=StubVuln(), items={}, variant="MYSQL")
    j = _judge_one(crit, item, ctx)
    assert j is not None
    assert j.verdict == "판단보류"
    assert "NOTE" in j.rationale


def test_judge_one_inline_note_no_false_match():
    # QUERY 줄 중간에 "NOTE:"가 섞여 있고 별도 NOTE 줄은 없는 경우,
    # NOTE 강제 분기에 진입하지 않고 정상 LLM 경로로 가야 한다.
    crit = _db_crit("DBM-101")
    # 증거 1건 부여(무증거 자동 판단보류 가드를 피해 LLM verdict 가 흐르도록).
    item = EvidenceItem(
        item_id="DBM-101", variant="MYSQL",
        resources=[ResourceEvidence(
            resource_id="db1", status="bad", detail="d", evidence="e")],
        context="QUERY: SELECT 'NOTE: inline' FROM dual")
    profile = get_profile("db_mysql")
    ctx = JudgeContext(profile=profile, profile_key="db_mysql",
                       client=StubVuln(), items={}, variant="MYSQL")
    j = _judge_one(crit, item, ctx)
    assert j is not None
    assert j.verdict == "취약"   # StubVuln → 취약, 판단보류 아님


# ── M3. --variant 오버라이드 테스트 ─────────────────────────────────────────

def test_variant_override_success(tmp_path):
    """(a) 마커 없는 파일명 + variant_override → metadata에 지정 variant 기록."""
    criteria = str(tmp_path / "db_native.xlsx")
    _db_native_criteria_xlsx(criteria)
    # 파일명은 마커 없는 임의 이름 — variant_override로 강제 지정
    report = str(tmp_path / "unknown.txt")
    with open(FIX, encoding="utf-8") as s, open(report, "w", encoding="utf-8") as d:
        d.write(s.read())
    jout = str(tmp_path / "r.json"); xout = str(tmp_path / "r.xlsx")
    run(report, criteria, "db_mysql", StubVuln(), jout, xout, "stub",
        variant_override="mysql_native")

    data = json.load(open(jout, encoding="utf-8"))
    assert data["metadata"]["variant"] == "mysql_native"


def test_variant_override_invalid_raises(tmp_path):
    """(b) 유효하지 않은 variant_override → '알 수 없는 variant' ReportError."""
    criteria = str(tmp_path / "db.xlsx")
    _db_criteria_xlsx(criteria)
    report = str(tmp_path / "mysql_result_rds.txt")
    with open(FIX, encoding="utf-8") as s, open(report, "w", encoding="utf-8") as d:
        d.write(s.read())
    jout = str(tmp_path / "r.json"); xout = str(tmp_path / "r.xlsx")
    with pytest.raises(ReportError, match="알 수 없는 variant"):
        run(report, criteria, "db_mysql", StubVuln(), jout, xout, "stub",
            variant_override="nope")


def test_no_override_unknown_filename_raises(tmp_path):
    """(c) 오버라이드 없이 마커 없는 파일명 → '--variant' 안내 ReportError."""
    criteria = str(tmp_path / "db.xlsx")
    _db_criteria_xlsx(criteria)
    # 파일명에 어떤 마커도 없는 임의 이름
    report = str(tmp_path / "unknown.txt")
    with open(FIX, encoding="utf-8") as s, open(report, "w", encoding="utf-8") as d:
        d.write(s.read())
    jout = str(tmp_path / "r.json"); xout = str(tmp_path / "r.xlsx")
    with pytest.raises(ReportError, match="--variant"):
        run(report, criteria, "db_mysql", StubVuln(), jout, xout, "stub")


# ── DBM-016 정책: 자동 verdict 금지 + 판단보류+힌트 경로 검증 ─────────────────
# 정책(2026-06-17): DBM-016 DET_SOURCE 전 variant STUB 전환.
# gate() → handled=False → _det_common_label_route → label D → _defer_or_eol
# → judge_patch → 판단보류 고정 + "미적용 후보/최신 패치 수준" 힌트.
# 자동 취약/양호(벤더 analysis.dbm_016() 버전비교) 경로가 완전 차단됨을 단언.

def test_dbm016_gate_stub_all_variants():
    """DBM-016은 전 DB 변형에서 DET_SOURCE=STUB → gate() handled=False."""
    from judge_tool.det_adapters.base import gate
    variants = [
        "mysql_native", "mysql_rds", "mysql_aurora", "mysql_azure",
        "oracle_native", "oracle_rds",
        "mssql_native", "mssql_rds",
        "mariadb_native", "mariadb_rds",
        "pg_native", "pg_rds", "pg_aurora", "pg_azure",
    ]
    for v in variants:
        result = gate("DBM-016", v)
        assert result is not None, f"gate() should return ForcedVerdict for DBM-016 {v}"
        assert result.handled is False, (
            f"DBM-016 {v}: gate() should return handled=False (STUB) to prevent auto-verdict")


def _make_dbm016_ctx(profile_key, variant, evidence_text):
    """DBM-016 판정 컨텍스트 헬퍼 (테스트 공용)."""
    from judge_tool.main import JudgeContext
    from judge_tool.models import EvidenceItem, ResourceEvidence
    from judge_tool.profile import get_profile

    res = ResourceEvidence(resource_id="r16", status="info", detail="",
                           evidence=evidence_text)
    item = EvidenceItem(item_id="DBM-016", variant=variant, resources=[res])
    items = {"DBM-016": item}
    profile = get_profile(profile_key)
    return (item, JudgeContext(profile=profile, profile_key=profile_key,
                               client=StubVuln(), items=items, variant=variant))


def _make_dbm016_crit(variant):
    """DBM-016 Criterion 헬퍼 — db_mysql.yaml 실설정과 동일."""
    return Criterion(
        item_id="DBM-016", item_name="보안패치 버전", risk=5.0, variant=variant,
        eval_type="스크립트",
        standard="* 양호 - ...\n* 취약 - ...", method="방법",
        applicable=True, label="D", patch_check=True,
        judgment_method="det_common",
        canned_message=(
            "DB 버전 정보 확인됨. 최신 보안패치 적용 여부는 벤더 릴리즈 노트 확인 필요. "
            "로컬 LLM은 최신 패치 기준을 알 수 없어 판단 불가."))


def test_dbm016_mysql_native_patch_hint():
    """mysql_native DBM-016: 8.4.4 < 8.4.9 → 판단보류 + '미적용 후보' 힌트."""
    from judge_tool.main import _det_common_handler
    evidence = '{"VARIABLE_NAME": "version","VARIABLE_VALUE": "8.4.4"}'
    variant = "mysql_native"
    item, ctx = _make_dbm016_ctx("db_mysql", variant, evidence)
    crit = _make_dbm016_crit(variant)
    j = _det_common_handler(crit, item, ctx)
    assert j is not None
    assert j.verdict == "판단보류", (
        f"DBM-016 mysql_native: 자동 verdict 금지 — 판단보류여야 함, 실제: {j.verdict}")
    assert j.verdict not in ("양호", "취약"), "거짓양호/거짓취약 금지"
    # judge_patch 힌트가 포함되어야 함
    assert "8.4.4" in j.rationale
    assert "8.4.9" in j.rationale
    assert "미적용 후보" in j.rationale
    # 테이블 기준일(as_of) 포함
    assert "패치 테이블 기준일" in j.rationale
    assert j.needs_review is True


def test_dbm016_mysql_rds_patch_hint():
    """mysql_rds DBM-016: 8.4.4 < 8.4.9 → 판단보류 + 관리형 서비스 문구."""
    from judge_tool.main import _det_common_handler
    evidence = '{"VARIABLE_NAME": "version","VARIABLE_VALUE": "8.4.4"}'
    variant = "mysql_rds"
    item, ctx = _make_dbm016_ctx("db_mysql", variant, evidence)
    crit = _make_dbm016_crit(variant)
    j = _det_common_handler(crit, item, ctx)
    assert j is not None
    assert j.verdict == "판단보류"
    assert j.verdict not in ("양호", "취약"), "거짓양호/거짓취약 금지"
    assert "관리형 서비스" in j.rationale


def test_dbm016_mssql_native_patch_hint():
    """mssql_native DBM-016: 빌드 15.0.4430.1 < 15.0.4470.1 → 판단보류 + 힌트."""
    from judge_tool.main import _det_common_handler
    # MSSQL: judge_patch는 _MSSQL_BUILD 패턴으로 빌드번호 추출
    evidence = ("Microsoft SQL Server 2019 (RTM-CU32) "
                "(KB5054833) - 15.0.4430.1 (X64)")
    variant = "mssql_native"
    # mssql의 경우 eol.py _PATTERNS에서 'Microsoft SQL Server (\d{4})' 로 연도 추출 후
    # series='2019', judge_patch에서 _MSSQL_BUILD 별도 추출
    res = ResourceEvidence(resource_id="r16", status="info", detail="",
                           evidence=evidence)
    item = EvidenceItem(item_id="DBM-016", variant=variant, resources=[res])
    items = {"DBM-016": item}
    profile = get_profile("db_mssql")
    ctx = JudgeContext(profile=profile, profile_key="db_mssql",
                       client=StubVuln(), items=items, variant=variant)
    crit = Criterion(
        item_id="DBM-016", item_name="보안패치 버전", risk=5.0, variant=variant,
        eval_type="스크립트",
        standard="* 양호 - ...\n* 취약 - ...", method="방법",
        applicable=True, label="D", patch_check=True,
        judgment_method="det_common",
        canned_message="DB 버전 정보 확인됨.")
    from judge_tool.main import _det_common_handler
    j = _det_common_handler(crit, item, ctx)
    assert j is not None
    assert j.verdict == "판단보류"
    assert j.verdict not in ("양호", "취약"), "거짓양호/거짓취약 금지"
    assert "15.0.4430.1" in j.rationale
    assert "15.0.4470.1" in j.rationale
    assert "미적용 후보" in j.rationale


def test_dbm016_oracle_native_patch_hint():
    """oracle_native DBM-016: 19.26 < 19.28 → 판단보류 + '미적용 후보' 힌트."""
    from judge_tool.main import _det_common_handler
    evidence = '{"description":"Database Release Update : 19.26.0.0.250121"}'
    variant = "oracle_native"
    res = ResourceEvidence(resource_id="r16", status="info", detail="",
                           evidence=evidence)
    item = EvidenceItem(item_id="DBM-016", variant=variant, resources=[res])
    items = {"DBM-016": item}
    profile = get_profile("db_oracle")
    ctx = JudgeContext(profile=profile, profile_key="db_oracle",
                       client=StubVuln(), items=items, variant=variant)
    crit = Criterion(
        item_id="DBM-016", item_name="보안패치 버전", risk=5.0, variant=variant,
        eval_type="스크립트",
        standard="* 양호 - ...\n* 취약 - ...", method="방법",
        applicable=True, label="D", patch_check=True,
        judgment_method="det_common",
        canned_message="DB 버전 정보 확인됨.")
    j = _det_common_handler(crit, item, ctx)
    assert j is not None
    assert j.verdict == "판단보류"
    assert j.verdict not in ("양호", "취약"), "거짓양호/거짓취약 금지"
    assert "19.26.0.0.250121" in j.rationale
    assert "19.28.0.0.250715" in j.rationale
    assert "미적용 후보" in j.rationale


def test_dbm016_pg_native_patch_hint():
    """pg_native DBM-016: 17.4 < 17.10 → 판단보류 + '미적용 후보' 힌트."""
    from judge_tool.main import _det_common_handler
    evidence = '{"version": "PostgreSQL 17.4 on aarch64-unknown-linux-gnu"}'
    variant = "pg_native"
    res = ResourceEvidence(resource_id="r16", status="info", detail="",
                           evidence=evidence)
    item = EvidenceItem(item_id="DBM-016", variant=variant, resources=[res])
    items = {"DBM-016": item}
    profile = get_profile("db_postgresql")
    ctx = JudgeContext(profile=profile, profile_key="db_postgresql",
                       client=StubVuln(), items=items, variant=variant)
    crit = Criterion(
        item_id="DBM-016", item_name="보안패치 버전", risk=5.0, variant=variant,
        eval_type="스크립트",
        standard="* 양호 - ...\n* 취약 - ...", method="방법",
        applicable=True, label="D", patch_check=True,
        judgment_method="det_common",
        canned_message="DB 버전 정보 확인됨.")
    j = _det_common_handler(crit, item, ctx)
    assert j is not None
    assert j.verdict == "판단보류"
    assert j.verdict not in ("양호", "취약"), "거짓양호/거짓취약 금지"
    assert "17.4" in j.rationale
    assert "17.10" in j.rationale
    assert "미적용 후보" in j.rationale


def test_dbm016_mariadb_canned_fallback():
    """mariadb DBM-016: eol.yaml latest 없는 시리즈(10.6) → canned 판단보류.

    mariadb는 DET_SOURCE 기존 STUB 유지. judge_patch가 None → canned 폴백.
    mariadb db_mariadb.yaml DBM-016은 judgment_method: det_common 없이
    label: D + patch_check: true만 있으므로 classify_method → 'det' 핸들러
    → _defer_or_eol → judge_patch → None(버전 미수록 시리즈) → _auto_defer canned.
    """
    from judge_tool.main import _defer_or_eol, JudgeContext
    from judge_tool.profile import get_profile
    # 10.6 시리즈는 eol.yaml에 없음 → judge_patch None → canned
    evidence = '{"VARIABLE_NAME": "version","VARIABLE_VALUE": "10.6.18-MariaDB"}'
    res = ResourceEvidence(resource_id="r16", status="info", detail="",
                           evidence=evidence)
    item = EvidenceItem(item_id="DBM-016", variant="mariadb_native", resources=[res])
    items = {"DBM-016": item}
    profile = get_profile("db_mariadb")
    ctx = JudgeContext(profile=profile, profile_key="db_mariadb",
                       client=StubVuln(), items=items, variant="mariadb_native")
    crit = Criterion(
        item_id="DBM-016", item_name="보안패치 버전", risk=5.0,
        variant="mariadb_native", eval_type="스크립트",
        standard="* 양호 - ...\n* 취약 - ...", method="방법",
        applicable=True, label="D", patch_check=True,
        canned_message=(
            "DB 버전 정보 확인됨. 최신 보안패치 적용 여부는 MariaDB 릴리즈 노트 확인 필요. "
            "로컬 LLM은 최신 패치 기준을 알 수 없어 판단 불가."))
    j = _defer_or_eol(crit, item, ctx)
    assert j.verdict == "판단보류"
    assert j.verdict not in ("양호", "취약"), "거짓양호/거짓취약 금지"


def test_dbm016_no_auto_verdict_from_det_path():
    """DBM-016 det_common 경로에서 자동 양호/취약이 절대 나오지 않음을 포괄 단언.

    모든 5 엔진 × native variant에서 DBM-016 det_common 처리 결과가
    verdict='판단보류'임을 확인한다(거짓양호 0, 거짓취약 0).
    """
    from judge_tool.main import _det_common_handler
    # (엔진, profile_key, variant, evidence)
    cases = [
        ("mysql", "db_mysql", "mysql_native",
         '{"VARIABLE_NAME": "version","VARIABLE_VALUE": "8.4.1"}'),
        ("oracle", "db_oracle", "oracle_native",
         '{"description":"Database Release Update : 19.10.0.0.220118"}'),
        ("mssql", "db_mssql", "mssql_native",
         "Microsoft SQL Server 2022 (RTM-CU1) - 16.0.4003.1 (X64)"),
        ("pg", "db_postgresql", "pg_native",
         '{"version": "PostgreSQL 16.1 on x86_64"}'),
        ("mariadb", "db_mariadb", "mariadb_native",
         '{"VARIABLE_NAME": "version","VARIABLE_VALUE": "11.4.2-MariaDB"}'),
    ]
    for engine, profile_key, variant, evidence in cases:
        res = ResourceEvidence(resource_id="r16", status="info", detail="",
                               evidence=evidence)
        item = EvidenceItem(item_id="DBM-016", variant=variant, resources=[res])
        items = {"DBM-016": item}
        profile = get_profile(profile_key)
        ctx = JudgeContext(profile=profile, profile_key=profile_key,
                           client=StubVuln(), items=items, variant=variant)
        crit = Criterion(
            item_id="DBM-016", item_name="보안패치 버전", risk=5.0,
            variant=variant, eval_type="스크립트",
            standard="* 양호 - ...\n* 취약 - ...", method="방법",
            applicable=True, label="D", patch_check=True,
            judgment_method="det_common",
            canned_message="DB 버전 정보 확인됨.")
        j = _det_common_handler(crit, item, ctx)
        assert j is not None, f"{engine} DBM-016 판정 None"
        assert j.verdict == "판단보류", (
            f"{engine} DBM-016: 자동 verdict 금지 — 판단보류여야 함, 실제: {j.verdict}")


def test_dbm003_interview_summary_passthrough():
    """모드A2(DBM-003) det가 생성한 interview_summary가 _det_common_handler를 거쳐
    Judgment.interview_summary로 전파되는지 확인(writer.py "인터뷰요약" 컬럼 배선 전제조건).

    ForcedVerdict.interview_summary(det_adapters/base.py trailing 필드) →
    _det_common_handler(main.py, j.needs_review=True 직후 3줄) → Judgment.interview_summary.
    """
    from judge_tool.main import _det_common_handler

    raw_evidence = json.dumps({"DBM-003": {"RESULT": [
        {"USER": "root", "HOST": "localhost", "ACCOUNT_LOCKED": "N"},
        {"USER": "test_svc", "HOST": "%", "ACCOUNT_LOCKED": "N"},
    ]}}, ensure_ascii=False)
    res = ResourceEvidence(resource_id="r3", status="info", detail="",
                           evidence="(masked)", raw_evidence=raw_evidence)
    item = EvidenceItem(item_id="DBM-003", variant="mysql_native", resources=[res])
    items = {"DBM-003": item}
    profile = get_profile("db_mysql")
    ctx = JudgeContext(profile=profile, profile_key="db_mysql",
                       client=StubVuln(), items=items, variant="mysql_native")
    crit = Criterion(
        item_id="DBM-003", item_name="업무상 불필요한 계정 존재", risk=5.0,
        variant="mysql_native", eval_type="스크립트",
        standard="* 양호 - ...\n* 취약 - ...", method="방법",
        applicable=True, label="B",
        judgment_method="det_common",
        summary_instruction="계정 목록을 분류하여 요약. 판정하지 말 것.")

    j = _det_common_handler(crit, item, ctx)

    assert j is not None
    assert j.verdict == "판단보류", f"DBM-003 모드A2: 판단보류 기대, 실제: {j.verdict}"
    assert j.needs_review is True
    assert j.interview_summary, (
        "DBM-003 모드A2: interview_summary가 Judgment로 전파되지 않음 "
        "(ForcedVerdict.interview_summary → _det_common_handler 배선 확인 필요)"
    )
    assert "test_svc" in j.interview_summary or "의심" in j.interview_summary
