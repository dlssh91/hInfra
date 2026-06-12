"""네트워크 장비 기준 로더 + main.run 통합 테스트 (작업④ — 합성 픽스처 기반).

주의: 합성 xlsx/XML 회귀 테스트이며, 실수집 데이터 대상 e2e는 샘플 미확보로
보류 상태다. results/ 실데이터는 사용하지 않는다.
"""
import json

import openpyxl
import pytest

from judge_tool.criteria_loader import load_criteria
from judge_tool.main import run
from judge_tool.profile import NETWORK

# 컬럼 배치: id=2, name=7, risk=8, standard=18, method=19, cisco_app=33
_CISCO_APP_COL = 33
_STD_COL = 18
_MTH_COL = 19


def _net_xlsx(path, rows):
    """합성 '네트워크 장비' 시트.

    rows: [(item_id, cisco_applicable: bool, standard: str)]
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "네트워크 장비"
    # 헤더 (row 4)
    ws.cell(4, 2, "평가항목ID")
    ws.cell(4, 7, "평가항목")
    ws.cell(4, 8, "위험도")
    ws.cell(4, _STD_COL, "판단기준")
    ws.cell(4, _MTH_COL, "판단방법")
    ws.cell(4, _CISCO_APP_COL, "평가대상(CISCO)")
    r = 5
    for item_id, cisco_app, standard in rows:
        ws.cell(r, 2, item_id)
        ws.cell(r, 7, f"{item_id}항목")
        ws.cell(r, 8, 5.0)
        ws.cell(r, _STD_COL, standard)
        ws.cell(r, _MTH_COL, "확인방법")
        if cisco_app:
            ws.cell(r, _CISCO_APP_COL, "o")
        r += 1
    wb.save(path)


# ── 기준 로더 단위 테스트 ──────────────────────────────────────────────────────────

def test_cisco_applicable_with_o(tmp_path):
    """cisco: c33='o' → applicable=True, is_judgeable=True, standard/method 채워짐."""
    p = str(tmp_path / "net.xlsx")
    _net_xlsx(p, [("NET-001", True, "* 양호 - ...\n* 취약 - ...")])
    crit = load_criteria(p, NETWORK, profile_key="network")
    c = crit[("NET-001", "cisco")]
    assert c.applicable is True
    assert c.is_judgeable is True
    assert "양호" in c.standard
    assert c.method == "확인방법"
    assert c.label == "A"
    assert c.judgment_method == "llm"


def test_cisco_not_applicable_without_o(tmp_path):
    """cisco: c33!='o' → applicable=False, is_judgeable=False."""
    p = str(tmp_path / "net2.xlsx")
    _net_xlsx(p, [("NET-002", False, "* 양호 - ...")])
    crit = load_criteria(p, NETWORK, profile_key="network")
    c = crit[("NET-002", "cisco")]
    assert c.applicable is False
    assert c.is_judgeable is False


def test_cisco_with_o_but_empty_standard_not_judgeable(tmp_path):
    """cisco: c33='o' 이지만 c18 빈칸 → is_judgeable=False."""
    p = str(tmp_path / "net3.xlsx")
    _net_xlsx(p, [("NET-003", True, "")])
    crit = load_criteria(p, NETWORK, profile_key="network")
    c = crit[("NET-003", "cisco")]
    assert c.applicable is True     # 'o' 로 applicable 결정
    assert c.is_judgeable is False  # standard 없으면 is_judgeable False


def test_generic_standard_nonempty_applicable(tmp_path):
    """generic: c18 비공백 → applicable=True, is_judgeable=True (c33 무관)."""
    p = str(tmp_path / "net4.xlsx")
    # cisco_app=False 이지만 generic은 standard로 결정
    _net_xlsx(p, [("NET-010", False, "* 양호 - 기준\n* 취약 - 위반")])
    crit = load_criteria(p, NETWORK, profile_key="network")
    g = crit[("NET-010", "generic")]
    assert g.applicable is True
    assert g.is_judgeable is True


def test_generic_empty_standard_not_applicable(tmp_path):
    """generic: c18 빈칸 → applicable=False, is_judgeable=False."""
    p = str(tmp_path / "net5.xlsx")
    _net_xlsx(p, [("NET-011", True, "")])   # cisco_app=True이지만 generic은 standard 봄
    crit = load_criteria(p, NETWORK, profile_key="network")
    g = crit[("NET-011", "generic")]
    assert g.applicable is False
    assert g.is_judgeable is False


# ── main.run 통합(가벼운 스텁) ─────────────────────────────────────────────────

class StubVuln:
    def chat(self, system, user):
        return ('{"verdict":"취약","confidence":0.8,'
                '"rationale":"테스트","cited_evidence":["x"]}')


def _net_xml_content(vendor="Cisco Systems", version="IOS 15.2",
                     model="Catalyst 3750", body_items=None):
    """합성 network XML 문자열."""
    if body_items is None:
        body_items = [("NET-001", "no ip http server\nno ip telnet")]
    dumps = ""
    for cid, output in body_items:
        dumps += f"""
<dump>
<items><id>{cid}</id></items>
<output><![CDATA[{output}]]></output>
</dump>
"""
    return f"""<?xml version="1.0" encoding="UTF-8"?>
<script>
<asset>
<hostname>router1</hostname>
<vendor>{vendor}</vendor>
<model>{model}</model>
<version>{version}</version>
</asset>
<results>{dumps}</results>
</script>
"""


def _write_net_xml(path, content):
    with open(path, "w", encoding="utf-8") as fh:
        fh.write(content)


def test_run_cisco_asset_variant(tmp_path):
    """Cisco vendor → metadata.variant == "cisco"."""
    criteria = str(tmp_path / "net.xlsx")
    _net_xlsx(criteria, [
        ("NET-001", True, "* 양호 - ...\n* 취약 - ..."),
        ("NET-002", True, "* 양호 - ..."),
    ])
    report = str(tmp_path / "router1.xml")
    _write_net_xml(report, _net_xml_content(
        vendor="Cisco Systems",
        body_items=[("NET-001", "no ip http server"),
                    ("NET-002", "   ")]))
    jout = str(tmp_path / "r.json")
    xout = str(tmp_path / "r.xlsx")
    cov = run(report, criteria, "network", StubVuln(), jout, xout, "stub")

    data = json.load(open(jout, encoding="utf-8"))
    assert data["metadata"]["variant"] == "cisco"
    assert data["metadata"]["profile"] == "network"
    by_id = {j["item_id"]: j for j in data["judgments"]}
    # NET-001: 증거 있음 → 스텁 취약 + flag_vulnerable_for_review → needs_review True
    assert by_id["NET-001"]["verdict"] == "취약"
    assert by_id["NET-001"]["needs_review"] is True
    assert by_id["NET-001"]["script_status"] is None   # status_available=False
    assert by_id["NET-001"]["agreement"] == "N/A"
    # NET-002: 공백 output → 증거 없음 → 판단보류 강제
    assert by_id["NET-002"]["verdict"] == "판단보류"
    assert cov["judged"] == 2 and cov["expected"] == 2
    # xlsx 산출 확인
    wb = openpyxl.load_workbook(xout)
    assert "판정결과" in wb.sheetnames


def test_run_juniper_asset_generic_variant(tmp_path):
    """Juniper vendor → detect_variant="generic", ReportError 발생 안 함."""
    criteria = str(tmp_path / "net.xlsx")
    _net_xlsx(criteria, [
        ("NET-001", False, "* 양호 - ...\n* 취약 - ..."),  # cisco_app=False
    ])
    report = str(tmp_path / "juniper.xml")
    _write_net_xml(report, _net_xml_content(
        vendor="Juniper Networks",
        version="JUNOS 22.4",   # 비-Cisco 버전 명시
        model="SRX-320",        # 비-Cisco 모델 명시 (기본 Catalyst 3750 방지)
        body_items=[("NET-001", "set system no-telnet")]))
    jout = str(tmp_path / "j.json")
    xout = str(tmp_path / "j.xlsx")
    # ReportError가 발생하지 않아야 함
    cov = run(report, criteria, "network", StubVuln(), jout, xout, "stub")

    data = json.load(open(jout, encoding="utf-8"))
    assert data["metadata"]["variant"] == "generic"
    # generic은 standard 비공백이므로 is_judgeable=True
    by_id = {j["item_id"]: j for j in data["judgments"]}
    assert "NET-001" in by_id
    assert cov["judged"] == 1


def test_run_variant_override_generic(tmp_path):
    """--variant generic override → cisco 장비도 generic으로 판정."""
    criteria = str(tmp_path / "net.xlsx")
    _net_xlsx(criteria, [
        ("NET-001", True, "* 양호 - ...\n* 취약 - ..."),
    ])
    report = str(tmp_path / "cisco.xml")
    _write_net_xml(report, _net_xml_content(
        vendor="Cisco Systems",
        body_items=[("NET-001", "no ip http server")]))
    jout = str(tmp_path / "g.json")
    run(report, criteria, "network", StubVuln(), jout, str(tmp_path / "g.xlsx"),
        "stub", variant_override="generic")
    data = json.load(open(jout, encoding="utf-8"))
    assert data["metadata"]["variant"] == "generic"


def test_run_empty_output_pending(tmp_path):
    """빈 output → 판단보류(증거 없음 보수적 처리)."""
    criteria = str(tmp_path / "net.xlsx")
    _net_xlsx(criteria, [("NET-001", True, "* 양호 - ...")])
    report = str(tmp_path / "empty_out.xml")
    _write_net_xml(report, _net_xml_content(
        body_items=[("NET-001", "   ")]))
    jout = str(tmp_path / "e.json")
    run(report, criteria, "network", StubVuln(), jout, str(tmp_path / "e.xlsx"), "stub")
    data = json.load(open(jout, encoding="utf-8"))
    by_id = {j["item_id"]: j for j in data["judgments"]}
    assert by_id["NET-001"]["verdict"] == "판단보류"


def test_run_unknown_dump_id_skipped(tmp_path):
    """기준에 없는 dump id는 자연 스킵 (coverage에 포함 안 됨)."""
    criteria = str(tmp_path / "net.xlsx")
    _net_xlsx(criteria, [("NET-001", True, "* 양호 - ...")])
    report = str(tmp_path / "unknown.xml")
    _write_net_xml(report, _net_xml_content(
        body_items=[
            ("NET-001", "config line"),
            ("UNKNOWN-999", "extra dump"),
        ]))
    jout = str(tmp_path / "u.json")
    cov = run(report, criteria, "network", StubVuln(), jout,
              str(tmp_path / "u.xlsx"), "stub")
    data = json.load(open(jout, encoding="utf-8"))
    by_id = {j["item_id"]: j for j in data["judgments"]}
    assert "NET-001" in by_id
    assert "UNKNOWN-999" not in by_id
    assert cov["expected"] == 1   # 기준에 NET-001만 있음


def test_run_all_judgments_needs_review_true(tmp_path):
    """네트워크 판정은 flag_vulnerable_for_review=True → 전 취약 판정 needs_review True."""
    criteria = str(tmp_path / "net.xlsx")
    _net_xlsx(criteria, [
        ("NET-001", True, "* 양호 - ..."),
        ("NET-002", True, "* 양호 - ..."),
    ])
    report = str(tmp_path / "nr.xml")
    _write_net_xml(report, _net_xml_content(
        body_items=[
            ("NET-001", "vulnerable config"),
            ("NET-002", "also vulnerable"),
        ]))
    jout = str(tmp_path / "nr.json")
    run(report, criteria, "network", StubVuln(), jout, str(tmp_path / "nr.xlsx"), "stub")
    data = json.load(open(jout, encoding="utf-8"))
    for j in data["judgments"]:
        if j["verdict"] == "취약":
            assert j["needs_review"] is True
