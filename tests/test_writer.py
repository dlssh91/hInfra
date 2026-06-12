import json
import os

import openpyxl

from judge_tool.models import Criterion, Judgment
from judge_tool.writer import build_coverage, write_json, write_excel


def _judgment(item_id="PISM-001", needs_review=False, cited_evidence=None,
              confidence=0.9):
    if cited_evidence is None:
        cited_evidence = ["b1"]
    return Judgment(item_id, "통신구간 암호화", "AWS", 5.0, "취약", confidence,
                    "근거", cited_evidence, "스크립트 전체", False, "bad", "일치",
                    needs_review)


def _criteria():
    # loader 동치: cloud applicable = ("스크립트" in eval_type) and eval_type != "N/A".
    return {
        ("PISM-001", "AWS"): Criterion("PISM-001", "암호화", 5.0, "AWS",
                                       "스크립트", "기준", "방법", applicable=True),
        ("PISM-005", "AWS"): Criterion("PISM-005", "퍼블릭", 5.0, "AWS",
                                       "스크립트", "기준", "방법", applicable=True),
        ("PISM-006", "AWS"): Criterion("PISM-006", "분리", 3.0, "AWS",
                                       "관리체계", "기준", "방법",
                                       applicable=False),  # 스크립트 아님
    }


def test_coverage_lists_missing_script_items():
    cov = build_coverage(_criteria(), [_judgment("PISM-001")], "AWS")
    assert cov["expected"] == 2          # PISM-001, 005 (006은 관리체계라 제외)
    assert cov["judged"] == 1
    assert cov["missing"] == ["PISM-005"]


def test_write_json(tmp_path):
    path = os.path.join(tmp_path, "out.json")
    meta = {"model": "test", "criteria_version": "제2026-1호"}
    cov = {"expected": 2, "judged": 1, "missing": ["PISM-005"]}
    write_json([_judgment()], meta, cov, path)
    data = json.load(open(path, encoding="utf-8"))
    assert data["metadata"]["model"] == "test"
    assert data["coverage"]["missing"] == ["PISM-005"]
    assert data["judgments"][0]["item_id"] == "PISM-001"


def test_write_excel(tmp_path):
    path = os.path.join(tmp_path, "out.xlsx")
    meta = {"model": "test", "criteria_version": "제2026-1호",
            "generated_at": "2026-06-01 10:00:00", "source_file": "x.xml",
            "source_sha256": "abc", "tool_version": "0.1.0"}
    cov = {"expected": 2, "judged": 1, "missing": ["PISM-005"]}
    write_excel([_judgment(needs_review=True)], meta, cov, path)
    wb = openpyxl.load_workbook(path)
    assert "판정결과" in wb.sheetnames
    ws = wb["판정결과"]
    header = [c.value for c in ws[1]]
    assert "항목ID" in header and "판정" in header and "재검토" in header


def _full_meta():
    return {"model": "test", "criteria_version": "제2026-1호",
            "generated_at": "2026-06-01 10:00:00", "source_file": "x.xml",
            "source_sha256": "abc", "tool_version": "0.1.0"}


def test_empty_judgments_coverage_and_files(tmp_path):
    cov = build_coverage(_criteria(), [], "AWS")
    assert cov["judged"] == 0
    assert cov["missing"] == ["PISM-001", "PISM-005"]

    meta = _full_meta()
    jpath = os.path.join(tmp_path, "out.json")
    xpath = os.path.join(tmp_path, "out.xlsx")
    write_json([], meta, cov, jpath)
    write_excel([], meta, cov, xpath)
    assert os.path.exists(jpath)
    assert os.path.exists(xpath)


def test_excel_review_fill_highlight(tmp_path):
    path = os.path.join(tmp_path, "out.xlsx")
    cov = {"expected": 2, "judged": 2, "missing": []}
    write_excel([_judgment("PISM-001", needs_review=True),
                 _judgment("PISM-005", needs_review=False)],
                _full_meta(), cov, path)
    wb = openpyxl.load_workbook(path)
    ws = wb["판정결과"]
    # row 2 = needs_review=True -> 연노랑 하이라이트
    assert "FFF2CC" in str(ws.cell(2, 1).fill.fgColor.rgb)
    # row 3 = needs_review=False -> 하이라이트 없음
    assert "FFF2CC" not in str(ws.cell(3, 1).fill.fgColor.rgb)


def test_excel_summary_sheet_contents(tmp_path):
    path = os.path.join(tmp_path, "out.xlsx")
    cov = {"expected": 2, "judged": 1, "missing": ["PISM-005"]}
    meta = {**_full_meta(), "profile": "network", "variant": "generic"}
    write_excel([_judgment(needs_review=True)], meta, cov, path)
    wb = openpyxl.load_workbook(path)
    ws = wb["요약"]
    summary = {ws.cell(r, 1).value: ws.cell(r, 2).value
               for r in range(1, ws.max_row + 1)}
    assert summary["사용 모델"] == "test"
    assert summary["대상 항목수"] == 2
    assert summary["재검토 필요수"] == 1
    # H: 프로파일·변형이 요약 시트에 노출 (운영자가 generic 판정 인지 가능)
    assert summary["프로파일"] == "network"
    assert summary["변형"] == "generic"


def test_excel_cited_evidence_join(tmp_path):
    path = os.path.join(tmp_path, "out.xlsx")
    cov = {"expected": 1, "judged": 1, "missing": []}
    write_excel([_judgment(cited_evidence=["b1", "b2"])], _full_meta(),
                cov, path)
    wb = openpyxl.load_workbook(path)
    ws = wb["판정결과"]
    idx = [c.value for c in ws[1]].index("인용증거") + 1
    assert ws.cell(2, idx).value == "b1 | b2"


def test_missing_sort_deterministic():
    criteria = {
        ("PISM-005", "AWS"): Criterion("PISM-005", "퍼블릭", 5.0, "AWS",
                                       "스크립트", "기준", "방법"),
        ("PISM-001", "AWS"): Criterion("PISM-001", "암호화", 5.0, "AWS",
                                       "스크립트", "기준", "방법"),
        ("PISM-003", "AWS"): Criterion("PISM-003", "기타", 5.0, "AWS",
                                       "스크립트", "기준", "방법"),
    }
    cov = build_coverage(criteria, [], "AWS")
    assert cov["missing"] == ["PISM-001", "PISM-003", "PISM-005"]


def test_excel_cited_evidence_none_regression(tmp_path):
    path = os.path.join(tmp_path, "out.xlsx")
    cov = {"expected": 1, "judged": 1, "missing": []}
    j = _judgment()
    j.cited_evidence = None
    write_excel([j], _full_meta(), cov, path)
    assert os.path.exists(path)


def test_excel_confidence_none_regression(tmp_path):
    path = os.path.join(tmp_path, "out.xlsx")
    cov = {"expected": 1, "judged": 1, "missing": []}
    j = _judgment()
    j.confidence = None
    write_excel([j], _full_meta(), cov, path)
    wb = openpyxl.load_workbook(path)
    ws = wb["판정결과"]
    idx = [c.value for c in ws[1]].index("확신도") + 1
    assert ws.cell(2, idx).value == 0.0


def test_write_makes_parent_dir(tmp_path):
    nested = os.path.join(tmp_path, "sub", "deep")
    jpath = os.path.join(nested, "out.json")
    xpath = os.path.join(nested, "out.xlsx")
    cov = {"expected": 0, "judged": 0, "missing": []}
    write_json([], _full_meta(), cov, jpath)
    write_excel([], _full_meta(), cov, xpath)
    assert os.path.exists(jpath)
    assert os.path.exists(xpath)
