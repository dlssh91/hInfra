"""ISS 비-FW 장비(iss_device 프로파일) 기준 로더 단위 테스트.

합성 xlsx 픽스처 기반 — 실수집 데이터 없이 실행한다.
"""
import openpyxl
import pytest

from judge_tool.criteria_loader import load_criteria
from judge_tool.profile import ISS_DEVICE

# 컬럼 배치 (1-indexed, openpyxl 기준)
_ID_COL    = 2
_NAME_COL  = 7
_RISK_COL  = 8
_FW_COL    = 12  # 평가대상(FW) — 이 테스트에서는 사용 안 함(iss 프로파일 전용)
_VPN_COL   = 13
_IDS_COL   = 14
_IPS_COL   = 15
_DDOS_COL  = 16
_WAF_COL   = 17
_STD_COL   = 18  # 판단기준 (전 장비 공유)
_MTH_COL   = 19  # 판단방법


def _iss_xlsx(tmp_path, rows):
    """합성 '정보보호시스템 장비' 시트 생성.

    rows: [(item_id, vpn_app, ids_app, ips_app, ddos_app, waf_app, standard)]
    각 *_app는 bool — True면 해당 열에 'o' 기재.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "정보보호시스템 장비"
    ws.cell(4, _ID_COL, "평가항목ID")
    ws.cell(4, _NAME_COL, "평가항목")
    ws.cell(4, _RISK_COL, "위험도")
    ws.cell(4, _VPN_COL, "평가대상(VPN)")
    ws.cell(4, _IDS_COL, "평가대상(IDS)")
    ws.cell(4, _IPS_COL, "평가대상(IPS)")
    ws.cell(4, _DDOS_COL, "평가대상(DDoS)")
    ws.cell(4, _WAF_COL, "평가대상(WAF)")
    ws.cell(4, _STD_COL, "판단기준")
    ws.cell(4, _MTH_COL, "판단방법")
    r = 5
    for item_id, vpn, ids_, ips_, ddos_, waf_, standard in rows:
        ws.cell(r, _ID_COL, item_id)
        ws.cell(r, _NAME_COL, f"{item_id}항목")
        ws.cell(r, _RISK_COL, 4.0)
        ws.cell(r, _STD_COL, standard)
        ws.cell(r, _MTH_COL, "확인방법")
        if vpn:
            ws.cell(r, _VPN_COL, "o")
        if ids_:
            ws.cell(r, _IDS_COL, "o")
        if ips_:
            ws.cell(r, _IPS_COL, "o")
        if ddos_:
            ws.cell(r, _DDOS_COL, "o")
        if waf_:
            ws.cell(r, _WAF_COL, "o")
        r += 1
    p = str(tmp_path / "iss_device.xlsx")
    wb.save(p)
    return p


# ── 적용성 매트릭스 ──────────────────────────────────────────────────────────

def test_iss001_applicable_all_variants(tmp_path):
    """ISS-001: 전 장비 'o' → vpn/ids/ips/ddos/waf 모두 applicable=True."""
    p = _iss_xlsx(tmp_path, [
        ("ISS-001", True, True, True, True, True, "* 양호 - ...\n* 취약 - ..."),
    ])
    crit = load_criteria(p, ISS_DEVICE, profile_key="iss_device")
    for variant in ("vpn", "ids", "ips", "ddos", "waf"):
        c = crit[("ISS-001", variant)]
        assert c.applicable is True, f"{variant} should be applicable"
        assert c.is_judgeable is True
        assert c.judgment_method == "llm"


def test_iss030_vpn_not_applicable(tmp_path):
    """ISS-030: VPN 열 None → vpn 변형 applicable=False 자동 제외."""
    p = _iss_xlsx(tmp_path, [
        # ISS-030: vpn=False (None), 나머지도 False — FW 전용 항목 재현
        ("ISS-030", False, False, False, False, False, "* 취약 - any-any 허용 정책"),
    ])
    crit = load_criteria(p, ISS_DEVICE, profile_key="iss_device")
    c = crit[("ISS-030", "vpn")]
    assert c.applicable is False
    assert c.is_judgeable is False


def test_iss030_generic_applicable_but_label_c(tmp_path):
    """ISS-030: generic은 applies_when_standard → applicable=True.
    iss_device.yaml label C → judgment_method='det', canned_message 포함.
    """
    p = _iss_xlsx(tmp_path, [
        ("ISS-030", False, False, False, False, False, "* 취약 - any-any 허용 정책"),
    ])
    crit = load_criteria(p, ISS_DEVICE, profile_key="iss_device")
    c = crit[("ISS-030", "generic")]
    assert c.applicable is True   # col18 비공백 → applies_when_standard 발동
    assert c.label == "C"
    assert c.judgment_method == "det"
    assert c.canned_message is not None
    assert "--profile iss" in c.canned_message


def test_iss003_vpn_only(tmp_path):
    """ISS-003(DMZ): VPN만 'o' → vpn=applicable, ids/ips/ddos/waf=not applicable."""
    p = _iss_xlsx(tmp_path, [
        ("ISS-003", True, False, False, False, False, "* 취약 - DMZ 미구성"),
    ])
    crit = load_criteria(p, ISS_DEVICE, profile_key="iss_device")
    assert crit[("ISS-003", "vpn")].applicable is True
    for variant in ("ids", "ips", "ddos", "waf"):
        assert crit[("ISS-003", variant)].applicable is False


def test_unregistered_item_defaults_label_a_llm(tmp_path):
    """iss_device.yaml 미등재 항목(ISS-010 등) → label A, judgment_method 'llm'."""
    p = _iss_xlsx(tmp_path, [
        ("ISS-010", True, True, True, True, True, "* 양호 - TCP/UDP 탐지 패턴 적용"),
    ])
    crit = load_criteria(p, ISS_DEVICE, profile_key="iss_device")
    c = crit[("ISS-010", "vpn")]
    assert c.label == "A"
    assert c.judgment_method == "llm"


def test_iss_device_profile_registered():
    """get_profile('iss_device') 성공: parser=iss_xml, generic.applies_when_standard=True."""
    from judge_tool.profile import get_profile
    prof = get_profile("iss_device")
    assert prof.parser == "iss_xml"
    assert "generic" in prof.variants
    assert prof.variants["generic"].applies_when_standard is True
    for variant in ("vpn", "ids", "ips", "ddos", "waf"):
        assert variant in prof.variants
