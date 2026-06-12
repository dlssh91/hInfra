"""ISS 기준 로더 + judgment_method 오버라이드 테스트 (작업②).

iss.yaml 판단방식 오버라이드(fw_policy/det)가 criteria_loader.load_criteria()에서
올바르게 반영되는지 합성 xlsx 픽스처로 확인한다. 실데이터 불필요.
"""
import openpyxl
import pytest

from judge_tool.criteria_loader import load_criteria
from judge_tool.profile import ISS

# ISS 프로파일 컬럼 배치
_ID_COL = 2
_NAME_COL = 7
_RISK_COL = 8
_APP_COL = 12    # fw variant applicability
_STD_COL = 18    # standard
_MTH_COL = 19    # method


def _iss_xlsx(path, rows):
    """합성 '정보보호시스템 장비' 시트.

    rows: [(item_id, fw_applicable: bool, standard: str, method: str)]
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "정보보호시스템 장비"
    # 헤더 (row 4)
    ws.cell(4, _ID_COL, "평가항목ID")
    ws.cell(4, _NAME_COL, "평가항목")
    ws.cell(4, _RISK_COL, "위험도")
    ws.cell(4, _APP_COL, "평가대상(FW)")
    ws.cell(4, _STD_COL, "판단기준")
    ws.cell(4, _MTH_COL, "판단방법")
    r = 5
    for item_id, fw_app, standard, method in rows:
        ws.cell(r, _ID_COL, item_id)
        ws.cell(r, _NAME_COL, f"{item_id} 항목명")
        ws.cell(r, _RISK_COL, 5.0)
        ws.cell(r, _STD_COL, standard)
        ws.cell(r, _MTH_COL, method)
        if fw_app:
            ws.cell(r, _APP_COL, "o")
        r += 1
    wb.save(str(path))


# ─ fw_policy judgment_method 오버라이드 (ISS-030~037, ISS-041) ─────────────────

def test_iss030_judgment_method_fw_policy(tmp_path):
    """iss.yaml에서 ISS-030: judgment_method=fw_policy → criteria에 반영."""
    p = str(tmp_path / "iss.xlsx")
    _iss_xlsx(p, [("ISS-030", True, "* 양호 - ...\n* 취약 - ...", "확인방법")])
    crit = load_criteria(p, ISS, profile_key="iss")
    c = crit[("ISS-030", "fw")]
    assert c.judgment_method == "fw_policy"
    assert c.label == "A"
    assert c.applicable is True


def test_iss031_judgment_method_fw_policy(tmp_path):
    p = str(tmp_path / "iss031.xlsx")
    _iss_xlsx(p, [("ISS-031", True, "기준텍스트", "방법")])
    crit = load_criteria(p, ISS, profile_key="iss")
    assert crit[("ISS-031", "fw")].judgment_method == "fw_policy"


def test_iss037_judgment_method_fw_policy(tmp_path):
    p = str(tmp_path / "iss037.xlsx")
    _iss_xlsx(p, [("ISS-037", True, "기준", "방법")])
    crit = load_criteria(p, ISS, profile_key="iss")
    assert crit[("ISS-037", "fw")].judgment_method == "fw_policy"


def test_iss041_judgment_method_fw_policy(tmp_path):
    p = str(tmp_path / "iss041.xlsx")
    _iss_xlsx(p, [("ISS-041", True, "기준", "방법")])
    crit = load_criteria(p, ISS, profile_key="iss")
    assert crit[("ISS-041", "fw")].judgment_method == "fw_policy"


# ─ label C → det (ISS-038/039/040) ───────────────────────────────────────────

def test_iss038_label_c_judgment_det(tmp_path):
    """iss.yaml에서 ISS-038: label=C → classify_method → 'det'."""
    p = str(tmp_path / "iss038.xlsx")
    _iss_xlsx(p, [("ISS-038", True, "기준", "방법")])
    crit = load_criteria(p, ISS, profile_key="iss")
    c = crit[("ISS-038", "fw")]
    assert c.label == "C"
    assert c.judgment_method == "det"
    assert c.canned_message is not None and len(c.canned_message) > 0


def test_iss039_label_c(tmp_path):
    p = str(tmp_path / "iss039.xlsx")
    _iss_xlsx(p, [("ISS-039", True, "기준", "방법")])
    crit = load_criteria(p, ISS, profile_key="iss")
    c = crit[("ISS-039", "fw")]
    assert c.label == "C"
    assert c.judgment_method == "det"


def test_iss040_label_c(tmp_path):
    p = str(tmp_path / "iss040.xlsx")
    _iss_xlsx(p, [("ISS-040", True, "기준", "방법")])
    crit = load_criteria(p, ISS, profile_key="iss")
    c = crit[("ISS-040", "fw")]
    assert c.label == "C"
    assert c.judgment_method == "det"


# ─ yaml_method 오버라이드가 classify_method를 이김 ───────────────────────────

def test_yaml_fw_policy_overrides_classify_method(tmp_path):
    """label=A → classify_method='llm' 이지만 yaml judgment_method='fw_policy'가 우선."""
    p = str(tmp_path / "override.xlsx")
    _iss_xlsx(p, [("ISS-030", True, "기준", "방법")])
    crit = load_criteria(p, ISS, profile_key="iss")
    c = crit[("ISS-030", "fw")]
    # classify_method(label="A", has_summary=False, in_empty_means_good=False) = "llm"
    # 하지만 yaml에서 fw_policy로 오버라이드
    assert c.judgment_method == "fw_policy"
    assert c.judgment_method != "llm"


# ─ yaml에 없는 항목 (ISS-001 등) → 기본 label A + llm ──────────────────────────

def test_iss001_not_in_yaml_defaults_llm(tmp_path):
    """ISS-001은 iss.yaml에 없어 기본 label=A, judgment_method=llm."""
    p = str(tmp_path / "iss001.xlsx")
    _iss_xlsx(p, [("ISS-001", True, "기준텍스트", "방법")])
    crit = load_criteria(p, ISS, profile_key="iss")
    c = crit[("ISS-001", "fw")]
    assert c.label == "A"
    assert c.judgment_method == "llm"


# ─ fw not applicable (applicability_col 없음) ────────────────────────────────

def test_fw_not_applicable_without_o(tmp_path):
    """fw applicability_col에 'o' 없으면 applicable=False."""
    p = str(tmp_path / "noapp.xlsx")
    _iss_xlsx(p, [("ISS-030", False, "기준", "방법")])
    crit = load_criteria(p, ISS, profile_key="iss")
    c = crit[("ISS-030", "fw")]
    assert c.applicable is False
    assert c.is_judgeable is False


# ─ 여러 항목 일괄 로드 ────────────────────────────────────────────────────────

def test_multiple_items_loaded(tmp_path):
    """여러 항목을 한 시트에 넣어도 모두 정상 로드된다."""
    p = str(tmp_path / "multi.xlsx")
    _iss_xlsx(p, [
        ("ISS-030", True, "기준A", "방법A"),
        ("ISS-031", True, "기준B", "방법B"),
        ("ISS-038", True, "기준C", "방법C"),
        ("ISS-001", True, "기준D", "방법D"),
    ])
    crit = load_criteria(p, ISS, profile_key="iss")
    assert crit[("ISS-030", "fw")].judgment_method == "fw_policy"
    assert crit[("ISS-031", "fw")].judgment_method == "fw_policy"
    assert crit[("ISS-038", "fw")].judgment_method == "det"
    assert crit[("ISS-001", "fw")].judgment_method == "llm"
