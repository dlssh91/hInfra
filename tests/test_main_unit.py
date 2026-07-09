import json
import os

import openpyxl
import pytest

import judge_tool.main as main_mod
from judge_tool.errors import ReportError
from judge_tool.main import main, run, _extract_criteria_version
from judge_tool.models import Criterion
from judge_tool.profile import CLOUD

FIXTURE_XML = os.path.join(
    os.path.dirname(__file__), "fixtures", "sample_aws_report.xml")
MALFORMED_XML = os.path.join(
    os.path.dirname(__file__), "fixtures", "malformed_report.xml")


class StubClient:
    """전건 고정 JSON 반환 대역 (Ollama 불필요)."""

    def __init__(self, verdict="양호", confidence=0.9):
        self._verdict = verdict
        self._confidence = confidence

    def chat(self, system, user):
        return (f'{{"verdict":"{self._verdict}","confidence":{self._confidence},'
                f'"rationale":"테스트 판정","cited_evidence":["x"]}}')


class BoomClient:
    """첫 호출만 예외를 던지고 이후는 정상 JSON을 반환하는 대역.

    부분 실패 격리(I-1) 검증용. 항목 단위 예외가 전체 run을 중단시키지
    않고 판단보류+needs_review로 격리되는지 확인한다.
    """

    def __init__(self):
        self.calls = 0

    def chat(self, system, user):
        self.calls += 1
        if self.calls == 1:
            raise RuntimeError("네트워크 폭발")
        return ('{"verdict":"양호","confidence":0.9,'
                '"rationale":"정상","cited_evidence":["x"]}')


class AllFailClient:
    """모든 chat 호출이 예외를 던지는 대역(전건 LLM 실패 검증용)."""

    def chat(self, system, user):
        raise RuntimeError("전건 네트워크 폭발")


_SECRET_MARKER = "SECRET-EVIDENCE-aws_access_key=AKIA0000"


class LeakyFailClient:
    """예외 메시지에 evidence 원문 같은 민감 문자열을 실어 던지는 대역.

    D-main 검증용: 예외 본문이 rationale/Excel '근거' 로 유출되면 안 된다.
    """

    def chat(self, system, user):
        raise RuntimeError(f"LLM 응답 파싱 실패: {_SECRET_MARKER}")


def _write_synthetic_criteria(path):
    """CLOUD 프로파일 포맷에 맞는 합성 평가기준 xlsx 생성.

    컬럼 레이아웃(CLOUD): id_col=2, name_col=6, risk_col=7,
    AWS variant: eval_type_col=11, standard_col=17, method_col=13.
    데이터 시작행=5. 시트명='클라우드 관리체계'.

    sample_aws_report.xml 의 항목과 매칭되도록 base id 를 직접 기입한다:
      - PISM-001 : 스크립트 (script status=bad)
      - PISM-037 : 스크립트 (split 037_1/037_2 병합)
      - PISM-007 : 스크립트 (script status=review)
      - PISM-099 : N/A      (스킵 대상)
      - PISM-098 : 관리체계 only (비스크립트 → 스킵 대상)
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = CLOUD.sheet_name

    rows = [
        # (item_id, name, risk, eval_type, method, standard)
        ("PISM-001", "통신구간 암호화", 5, "스크립트", "방법1", "양호 기준 텍스트"),
        ("PISM-037", "비밀번호 정책", 3, "스크립트", "방법37", "양호 기준 텍스트"),
        ("PISM-007", "네트워크 접근제어", 4, "스크립트", "방법7", "양호 기준 텍스트"),
        ("PISM-099", "관리체계 항목", 2, "N/A", "방법99", "기준99"),
        ("PISM-098", "관리체계만", 2, "관리체계", "방법98", "기준98"),
    ]
    for i, (iid, name, risk, etype, method, standard) in enumerate(rows):
        r = CLOUD.data_start_row + i
        ws.cell(r, CLOUD.id_col, iid)
        ws.cell(r, CLOUD.name_col, name)
        ws.cell(r, CLOUD.risk_col, risk)
        aws = CLOUD.variants["AWS"]
        ws.cell(r, aws.eval_type_col, etype)
        ws.cell(r, aws.method_col, method)
        ws.cell(r, aws.standard_col, standard)
        # Azure variant 도 채워 둠(로더가 양 variant 로드)
        az = CLOUD.variants["Azure"]
        ws.cell(r, az.eval_type_col, etype)
        ws.cell(r, az.method_col, method)
        ws.cell(r, az.standard_col, standard)
    wb.save(path)


def _copy_fixture_xml(tmp_path, name="aws_report_synth.xml"):
    """fixture XML 을 variant 식별 가능한 파일명으로 tmp_path 에 복사."""
    dest = os.path.join(str(tmp_path), name)
    with open(FIXTURE_XML, encoding="utf-8") as src:
        content = src.read()
    with open(dest, "w", encoding="utf-8") as fh:
        fh.write(content)
    return dest


# --------------------------------------------------------------------------
# I-2: main() CLI 경로
# --------------------------------------------------------------------------

def test_main_cli_writes_outputs(tmp_path, monkeypatch, capsys,
                                 aws_report_path, criteria_xlsx_path):
    if not (os.path.exists(aws_report_path)
            and os.path.exists(criteria_xlsx_path)):
        pytest.skip("실제 보고서/평가기준 파일이 없어 CLI E2E 스킵")

    monkeypatch.setattr(main_mod, "OllamaClient",
                        lambda *a, **k: StubClient(verdict="취약"))

    main(["--report", aws_report_path,
          "--criteria", criteria_xlsx_path,
          "--out-dir", str(tmp_path),
          "--model", "stub"])

    base = os.path.splitext(os.path.basename(aws_report_path))[0]
    json_out = os.path.join(str(tmp_path), f"result_{base}.json")
    xlsx_out = os.path.join(str(tmp_path), f"result_{base}.xlsx")
    assert os.path.exists(json_out)
    assert os.path.exists(xlsx_out)

    out = capsys.readouterr().out
    assert "판정" in out


def test_main_cli_missing_required_arg():
    with pytest.raises(SystemExit):
        main(["--report", "x"])  # --criteria 누락 → argparse SystemExit


# --------------------------------------------------------------------------
# B′-3b: --aux-objects CLI 배선 (FW ObjectTable 실치환) E2E
# --------------------------------------------------------------------------

_KRFW_AUX_HEADER = ["룰 NUM", "출발지", "목적지", "서비스", "Protocol", "inbound",
                    "시간", "정책", "로그", "session-limit", "tcp", "활성화", "설명"]


def _write_krfw_named_dst_xlsx(path):
    """ISS-030(any-any) 판정용: dst가 named 그룹(ANY_GRP)인 단일 허용정책."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Detail"
    rows = [
        _KRFW_AUX_HEADER,
        ["1", "any", "ANY_GRP", None, "tcp", None, None,
         "허용", "Enable", None, None, "Enable", "테스트"],
    ]
    for r_idx, row in enumerate(rows, start=1):
        for c_idx, val in enumerate(row, start=1):
            ws.cell(r_idx, c_idx, val)
    wb.save(str(path))


def _write_iss030_criteria_xlsx(path):
    """ISS-030 단일 항목짜리 합성 '정보보호시스템 장비' 평가기준 xlsx."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "정보보호시스템 장비"
    ws.cell(4, 2, "평가항목ID")
    ws.cell(4, 7, "평가항목")
    ws.cell(4, 8, "위험도")
    ws.cell(4, 12, "평가대상(FW)")
    ws.cell(4, 18, "판단기준")
    ws.cell(4, 19, "판단방법")
    ws.cell(5, 2, "ISS-030")
    ws.cell(5, 7, "ISS-030 항목명")
    ws.cell(5, 8, 5.0)
    ws.cell(5, 12, "o")
    ws.cell(5, 18, "* 양호 - any-any 허용정책 없음\n* 취약 - 존재")
    ws.cell(5, 19, "정책 확인")
    wb.save(str(path))


def test_main_cli_parses_aux_objects_arg_and_threads_to_run(
    tmp_path, tmp_path_factory, monkeypatch
):
    """--aux-objects 값이 argparse dest(aux_objects)를 거쳐 run()의
    aux_objects_path 키워드 인자로 정확히 전달되는지 확인(배선 검증)."""
    report_path = str(tmp_path / "P99_정책.xlsx")
    _write_krfw_named_dst_xlsx(report_path)
    criteria_path = str(tmp_path / "기준.xlsx")
    _write_iss030_criteria_xlsx(criteria_path)
    aux_path = str(tmp_path / "objects.yaml")
    with open(aux_path, "w", encoding="utf-8") as f:
        f.write("address:\n  ANY_GRP: [\"0.0.0.0/0\"]\n")
    out_dir = tmp_path_factory.mktemp("aux_wire_out")

    captured = {}

    def _fake_run(*args, **kwargs):
        captured.update(kwargs)
        return {"judged": 0, "expected": 0, "missing": []}

    monkeypatch.setattr(main_mod, "run", _fake_run)
    main(["--report", report_path, "--criteria", criteria_path,
          "--profile", "iss", "--out-dir", str(out_dir),
          "--skip-preflight", "--aux-objects", aux_path])
    assert captured.get("aux_objects_path") == aux_path


def test_main_cli_aux_objects_omitted_threads_none(
    tmp_path, tmp_path_factory, monkeypatch
):
    """--aux-objects 미지정 시 run()에 aux_objects_path=None이 전달된다."""
    report_path = str(tmp_path / "P99_정책.xlsx")
    _write_krfw_named_dst_xlsx(report_path)
    criteria_path = str(tmp_path / "기준.xlsx")
    _write_iss030_criteria_xlsx(criteria_path)
    out_dir = tmp_path_factory.mktemp("aux_wire_out_none")

    captured = {}

    def _fake_run(*args, **kwargs):
        captured.update(kwargs)
        return {"judged": 0, "expected": 0, "missing": []}

    monkeypatch.setattr(main_mod, "run", _fake_run)
    main(["--report", report_path, "--criteria", criteria_path,
          "--profile", "iss", "--out-dir", str(out_dir),
          "--skip-preflight"])
    assert captured.get("aux_objects_path") is None


def test_fw_aux_objects_resolves_named_group_to_violation(
    tmp_path, tmp_path_factory
):
    """--aux-objects 주입 시 named dst 그룹이 0.0.0.0/0으로 실치환되어
    ISS-030(any-any 허용)이 '취약'로 실판정된다(미주입 시엔 판단보류)."""
    report_path = str(tmp_path / "P99_정책.xlsx")
    _write_krfw_named_dst_xlsx(report_path)
    criteria_path = str(tmp_path / "기준.xlsx")
    _write_iss030_criteria_xlsx(criteria_path)
    aux_path = str(tmp_path / "objects.yaml")
    with open(aux_path, "w", encoding="utf-8") as f:
        f.write("address:\n  ANY_GRP: [\"0.0.0.0/0\"]\n")

    out_dir_no_aux = tmp_path_factory.mktemp("out_no_aux")
    main(["--report", report_path, "--criteria", criteria_path,
          "--profile", "iss", "--out-dir", str(out_dir_no_aux),
          "--skip-preflight"])
    json_no_aux = out_dir_no_aux / "result_P99_정책.json"
    with open(json_no_aux, encoding="utf-8") as f:
        payload_no_aux = json.load(f)
    verdict_no_aux = next(
        j["verdict"] for j in payload_no_aux["judgments"]
        if j["item_id"] == "ISS-030")
    assert verdict_no_aux == "판단보류"

    out_dir_aux = tmp_path_factory.mktemp("out_aux")
    main(["--report", report_path, "--criteria", criteria_path,
          "--profile", "iss", "--out-dir", str(out_dir_aux),
          "--skip-preflight", "--aux-objects", aux_path])
    json_aux = out_dir_aux / "result_P99_정책.json"
    with open(json_aux, encoding="utf-8") as f:
        payload_aux = json.load(f)
    verdict_aux = next(
        j["verdict"] for j in payload_aux["judgments"]
        if j["item_id"] == "ISS-030")
    assert verdict_aux == "취약"


def test_fw_aux_objects_missing_file_raises_report_error(
    tmp_path, tmp_path_factory
):
    """--aux-objects에 존재하지 않는 경로를 주면 ReportError로 깔끔히 안내
    (fail-closed — 조용히 무시하고 기존 동작으로 폴백하지 않음)."""
    report_path = str(tmp_path / "P99_정책.xlsx")
    _write_krfw_named_dst_xlsx(report_path)
    criteria_path = str(tmp_path / "기준.xlsx")
    _write_iss030_criteria_xlsx(criteria_path)
    out_dir = tmp_path_factory.mktemp("out_missing_aux")

    with pytest.raises(SystemExit):
        main(["--report", report_path, "--criteria", criteria_path,
              "--profile", "iss", "--out-dir", str(out_dir),
              "--skip-preflight",
              "--aux-objects", str(tmp_path / "no_such_objects.yaml")])


def test_run_aux_objects_path_none_default_noop():
    """run()의 aux_objects_path 기본값 None → 시그니처 호출부(배치/단일 CLI)가
    인자를 생략해도 기존 호출과 동일하게 동작(키워드 인자 하위호환)."""
    import inspect
    sig = inspect.signature(run)
    assert sig.parameters["aux_objects_path"].default is None


# --------------------------------------------------------------------------
# I-3: CI 독립 결정적 run() (합성 입력)
# --------------------------------------------------------------------------

def test_run_synthetic_disagreement_needs_review(tmp_path):
    """StubClient 가 전건 '양호' → script status=bad/review 와 불일치 →
    needs_review True. N/A·관리체계 항목은 judged 에서 제외(스킵)."""
    report = _copy_fixture_xml(tmp_path)
    criteria = os.path.join(str(tmp_path), "criteria.xlsx")
    _write_synthetic_criteria(criteria)
    json_out = os.path.join(str(tmp_path), "result.json")
    xlsx_out = os.path.join(str(tmp_path), "result.xlsx")

    cov = run(report_path=report, criteria_path=criteria, profile_key="cloud",
              client=StubClient(verdict="양호", confidence=0.9),
              json_out=json_out, xlsx_out=xlsx_out, model_name="stub")

    data = json.load(open(json_out, encoding="utf-8"))
    by_id = {j["item_id"]: j for j in data["judgments"]}

    # 스크립트 항목만 판정됨
    assert set(by_id) == {"PISM-001", "PISM-037", "PISM-007"}
    # N/A / 관리체계 only 는 스킵
    assert "PISM-099" not in by_id
    assert "PISM-098" not in by_id

    # PISM-001: script bad vs llm 양호 → 불일치 → needs_review
    assert by_id["PISM-001"]["agreement"] == "불일치"
    assert by_id["PISM-001"]["needs_review"] is True

    # 분할항목 PISM-037: 037_1(bad)+037_2(good) overall=bad → 불일치
    assert by_id["PISM-037"]["agreement"] == "불일치"
    assert by_id["PISM-037"]["needs_review"] is True

    # coverage: expected=3(스크립트만), judged=3
    assert cov["expected"] == 3
    assert cov["judged"] == 3


def test_run_variant_identification_failure(tmp_path):
    """파일명에서 variant 를 식별 못 하면 ValueError."""
    bad_name = os.path.join(str(tmp_path), "unknown.xml")
    with open(bad_name, "w", encoding="utf-8") as fh:
        fh.write("<AuditReport/>")
    criteria = os.path.join(str(tmp_path), "criteria.xlsx")
    _write_synthetic_criteria(criteria)

    with pytest.raises(ValueError):
        run(report_path=bad_name, criteria_path=criteria, profile_key="cloud",
            client=StubClient(), json_out=os.path.join(str(tmp_path), "j.json"),
            xlsx_out=os.path.join(str(tmp_path), "x.xlsx"), model_name="stub")


def test_run_variant_identification_failure_is_report_error(tmp_path):
    """variant 식별 실패는 전용 ReportError(ValueError 하위)로 raise."""
    from judge_tool.errors import ReportError

    bad_name = os.path.join(str(tmp_path), "unknown.xml")
    with open(bad_name, "w", encoding="utf-8") as fh:
        fh.write("<AuditReport/>")
    criteria = os.path.join(str(tmp_path), "criteria.xlsx")
    _write_synthetic_criteria(criteria)

    with pytest.raises(ReportError):
        run(report_path=bad_name, criteria_path=criteria, profile_key="cloud",
            client=StubClient(), json_out=os.path.join(str(tmp_path), "j.json"),
            xlsx_out=os.path.join(str(tmp_path), "x.xlsx"), model_name="stub")


def test_main_propagates_unexpected_value_error(tmp_path, monkeypatch):
    """run 내부 콜트리의 우발적(비-ReportError) ValueError 는 main()이 삼키지
    않고 그대로 전파한다(SystemExit 로 가려지지 않아 디버깅 가능)."""
    from judge_tool.errors import ReportError

    rep_dir = tmp_path / "rep"
    crit_dir = tmp_path / "crit"
    out_dir = tmp_path / "out"
    for d in (rep_dir, crit_dir, out_dir):
        os.makedirs(str(d), exist_ok=True)
    report = _copy_fixture_xml(rep_dir)
    criteria = os.path.join(str(crit_dir), "criteria.xlsx")
    _write_synthetic_criteria(criteria)

    def _boom(*a, **k):
        raise ValueError("boom")

    monkeypatch.setattr(main_mod, "load_criteria", _boom)
    monkeypatch.setattr(main_mod, "OllamaClient",
                        lambda *a, **k: StubClient())

    with pytest.raises(ValueError) as ei:
        main(["--report", report, "--criteria", criteria,
              "--out-dir", str(out_dir), "--model", "stub"])
    # 전파된 예외가 ReportError 가 아니어야(우발적 버그) 한다.
    assert not isinstance(ei.value, ReportError)


def test_run_isolates_per_item_failure(tmp_path):
    """한 항목 judge 예외가 전체 run 을 중단시키지 않고, 실패 항목은
    판단보류+needs_review 로 격리된 채 출력이 생성된다(I-1)."""
    report = _copy_fixture_xml(tmp_path)
    criteria = os.path.join(str(tmp_path), "criteria.xlsx")
    _write_synthetic_criteria(criteria)
    json_out = os.path.join(str(tmp_path), "result.json")
    xlsx_out = os.path.join(str(tmp_path), "result.xlsx")

    cov = run(report_path=report, criteria_path=criteria, profile_key="cloud",
              client=BoomClient(), json_out=json_out, xlsx_out=xlsx_out,
              model_name="stub", skip_preflight=True)

    assert os.path.exists(json_out) and os.path.exists(xlsx_out)
    data = json.load(open(json_out, encoding="utf-8"))
    by_id = {j["item_id"]: j for j in data["judgments"]}

    # 3개 스크립트 항목 모두 결과에 존재(실패 항목도 격리되어 포함)
    assert set(by_id) == {"PISM-001", "PISM-037", "PISM-007"}
    # 정확히 한 항목이 예외로 판단보류 격리됨
    held = [j for j in data["judgments"] if j["verdict"] == "판단보류"]
    assert len(held) == 1
    assert held[0]["needs_review"] is True
    assert cov["judged"] == 3


# --------------------------------------------------------------------------
# I-2 (spec 6.4): 빈 판단기준(standard="") 항목 스킵
# --------------------------------------------------------------------------

def _write_empty_standard_criteria(path):
    """PISM-001 의 standard 를 비워 둔 합성 평가기준 xlsx.

    빈 판단기준은 LLM 호출 없이 스킵되어야 한다(spec 6.4).
    PISM-007 은 정상 기준을 둬 대조군으로 사용.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = CLOUD.sheet_name
    rows = [
        # (item_id, name, risk, eval_type, method, standard)
        ("PISM-001", "통신구간 암호화", 5, "스크립트", "방법1", ""),     # 빈 기준 → 스킵
        ("PISM-007", "네트워크 접근제어", 4, "스크립트", "방법7", "양호 기준"),
    ]
    for i, (iid, name, risk, etype, method, standard) in enumerate(rows):
        r = CLOUD.data_start_row + i
        ws.cell(r, CLOUD.id_col, iid)
        ws.cell(r, CLOUD.name_col, name)
        ws.cell(r, CLOUD.risk_col, risk)
        for vname in ("AWS", "Azure"):
            v = CLOUD.variants[vname]
            ws.cell(r, v.eval_type_col, etype)
            ws.cell(r, v.method_col, method)
            ws.cell(r, v.standard_col, standard)
    wb.save(path)


def test_run_skips_empty_standard_item(tmp_path):
    """standard 가 비어 있는 스크립트 항목은 judged 에서 제외(스킵)된다."""
    report = _copy_fixture_xml(tmp_path)
    criteria = os.path.join(str(tmp_path), "criteria.xlsx")
    _write_empty_standard_criteria(criteria)
    json_out = os.path.join(str(tmp_path), "result.json")
    xlsx_out = os.path.join(str(tmp_path), "result.xlsx")

    cov = run(report_path=report, criteria_path=criteria, profile_key="cloud",
              client=StubClient(verdict="취약"), json_out=json_out,
              xlsx_out=xlsx_out, model_name="stub")

    data = json.load(open(json_out, encoding="utf-8"))
    by_id = {j["item_id"]: j for j in data["judgments"]}
    # 빈 기준 PISM-001 은 스킵, PISM-007 만 판정
    assert "PISM-001" not in by_id
    assert "PISM-007" in by_id
    assert cov["judged"] == 1


# --------------------------------------------------------------------------
# 견고성: 전건 LLM 실패 / 빈 입력
# --------------------------------------------------------------------------

def test_run_all_llm_failures_isolated(tmp_path):
    """모든 chat 호출이 예외여도 run 은 예외 없이 완료되고, 모든 판정이
    판단보류 & needs_review 로 격리된다(부분 실패 격리가 전건에도 동작)."""
    report = _copy_fixture_xml(tmp_path)
    criteria = os.path.join(str(tmp_path), "criteria.xlsx")
    _write_synthetic_criteria(criteria)
    json_out = os.path.join(str(tmp_path), "result.json")
    xlsx_out = os.path.join(str(tmp_path), "result.xlsx")

    cov = run(report_path=report, criteria_path=criteria, profile_key="cloud",
              client=AllFailClient(), json_out=json_out, xlsx_out=xlsx_out,
              model_name="stub")

    assert os.path.exists(json_out) and os.path.exists(xlsx_out)
    data = json.load(open(json_out, encoding="utf-8"))
    judgments = data["judgments"]
    assert len(judgments) == 3
    assert cov["judged"] == 3
    for j in judgments:
        assert j["verdict"] == "판단보류"
        assert j["needs_review"] is True


def test_run_redacts_exception_body_from_rationale(tmp_path):
    """judge 예외 메시지(민감 evidence 포함)가 rationale 으로 유출되지 않는다.

    rationale 은 예외 타입명/고정문구만 담아야 한다(D-main).
    """
    report = _copy_fixture_xml(tmp_path)
    criteria = os.path.join(str(tmp_path), "criteria.xlsx")
    _write_synthetic_criteria(criteria)
    json_out = os.path.join(str(tmp_path), "result.json")
    xlsx_out = os.path.join(str(tmp_path), "result.xlsx")

    run(report_path=report, criteria_path=criteria, profile_key="cloud",
        client=LeakyFailClient(), json_out=json_out, xlsx_out=xlsx_out,
        model_name="stub")

    data = json.load(open(json_out, encoding="utf-8"))
    for j in data["judgments"]:
        assert j["verdict"] == "판단보류"
        # 예외 raw 본문(민감 문자열)이 rationale 에 새어나오면 안 됨
        assert _SECRET_MARKER not in j["rationale"]
        # 예외 타입명(고정문구)은 허용
        assert "RuntimeError" in j["rationale"]


# --------------------------------------------------------------------------
# E: 출력 디렉터리가 입력 데이터 디렉터리로 가는 것 차단(main CLI 경계)
# --------------------------------------------------------------------------

def test_main_rejects_out_dir_equal_to_report_dir(tmp_path, monkeypatch):
    report = _copy_fixture_xml(tmp_path)
    criteria = os.path.join(str(tmp_path / "crit"), "criteria.xlsx")
    os.makedirs(os.path.dirname(criteria), exist_ok=True)
    _write_synthetic_criteria(criteria)

    monkeypatch.setattr(main_mod, "OllamaClient",
                        lambda *a, **k: StubClient())

    # --out-dir 이 --report 가 위치한 디렉터리(tmp_path)와 동일 → 거부
    with pytest.raises((SystemExit, ValueError)):
        main(["--report", report, "--criteria", criteria,
              "--out-dir", str(tmp_path), "--model", "stub"])


def test_main_rejects_out_dir_equal_to_criteria_dir(tmp_path, monkeypatch):
    rep_dir = tmp_path / "rep"
    crit_dir = tmp_path / "crit"
    os.makedirs(str(rep_dir), exist_ok=True)
    os.makedirs(str(crit_dir), exist_ok=True)
    report = _copy_fixture_xml(rep_dir)
    criteria = os.path.join(str(crit_dir), "criteria.xlsx")
    _write_synthetic_criteria(criteria)

    monkeypatch.setattr(main_mod, "OllamaClient",
                        lambda *a, **k: StubClient())

    # --out-dir 이 --criteria 디렉터리와 동일 → 거부
    with pytest.raises((SystemExit, ValueError)):
        main(["--report", report, "--criteria", criteria,
              "--out-dir", str(crit_dir), "--model", "stub"])


def test_main_rejects_out_dir_under_report_dir(tmp_path, monkeypatch):
    rep_dir = tmp_path / "rep"
    crit_dir = tmp_path / "crit"
    os.makedirs(str(rep_dir), exist_ok=True)
    os.makedirs(str(crit_dir), exist_ok=True)
    report = _copy_fixture_xml(rep_dir)
    criteria = os.path.join(str(crit_dir), "criteria.xlsx")
    _write_synthetic_criteria(criteria)
    sub = os.path.join(str(rep_dir), "out")  # 입력 디렉터리 하위

    monkeypatch.setattr(main_mod, "OllamaClient",
                        lambda *a, **k: StubClient())

    with pytest.raises((SystemExit, ValueError)):
        main(["--report", report, "--criteria", criteria,
              "--out-dir", sub, "--model", "stub"])


def test_main_allows_separate_out_dir(tmp_path, monkeypatch, capsys):
    rep_dir = tmp_path / "rep"
    crit_dir = tmp_path / "crit"
    out_dir = tmp_path / "out"
    for d in (rep_dir, crit_dir, out_dir):
        os.makedirs(str(d), exist_ok=True)
    report = _copy_fixture_xml(rep_dir)
    criteria = os.path.join(str(crit_dir), "criteria.xlsx")
    _write_synthetic_criteria(criteria)

    monkeypatch.setattr(main_mod, "OllamaClient",
                        lambda *a, **k: StubClient(verdict="취약"))

    # 입력과 무관한 별도 디렉터리 → 정상 동작
    main(["--report", report, "--criteria", criteria,
          "--out-dir", str(out_dir), "--model", "stub"])

    base = os.path.splitext(os.path.basename(report))[0]
    assert os.path.exists(os.path.join(str(out_dir), f"result_{base}.json"))
    assert os.path.exists(os.path.join(str(out_dir), f"result_{base}.xlsx"))


# --------------------------------------------------------------------------
# 항목2: 손상 XML → main()이 트레이스백 대신 깔끔한 SystemExit
# --------------------------------------------------------------------------

def test_main_malformed_xml_exits_cleanly(tmp_path, monkeypatch, capsys):
    """손상 XML 입력 시 main()은 raw 트레이스백이 아니라 SystemExit으로
    종료하고 stderr에 명확한 에러 메시지를 출력한다."""
    crit_dir = tmp_path / "crit"
    out_dir = tmp_path / "out"
    rep_dir = tmp_path / "rep"
    for d in (crit_dir, out_dir, rep_dir):
        os.makedirs(str(d), exist_ok=True)
    # variant 식별 가능한 파일명으로 손상 XML 복사
    report = os.path.join(str(rep_dir), "aws_report_malformed.xml")
    with open(MALFORMED_XML, encoding="utf-8") as src:
        content = src.read()
    with open(report, "w", encoding="utf-8") as fh:
        fh.write(content)
    criteria = os.path.join(str(crit_dir), "criteria.xlsx")
    _write_synthetic_criteria(criteria)

    monkeypatch.setattr(main_mod, "OllamaClient",
                        lambda *a, **k: StubClient())

    with pytest.raises(SystemExit):
        main(["--report", report, "--criteria", criteria,
              "--out-dir", str(out_dir), "--model", "stub"])

    err = capsys.readouterr().err
    assert "파싱 실패" in err


# --------------------------------------------------------------------------
# 항목3: criteria_version 파일명에서 추출
# --------------------------------------------------------------------------

def test_extract_criteria_version_from_real_filename():
    p = ("ref/전자금융기반시설 보안 취약점 평가기준(제2026-1호) "
         "평가자용_2603개정.xlsx")
    assert _extract_criteria_version(p) == "제2026-1호"


def test_extract_criteria_version_fallback_when_no_pattern():
    assert _extract_criteria_version("criteria.xlsx") == "제2026-1호"


# --------------------------------------------------------------------------
# 스마트 런처: _needs_llm — Ollama 페일패스트 헬스체크 게이트 로직
# --------------------------------------------------------------------------

def _crit(item_id="X-001", variant="AWS", judgment_method="llm",
         summary_instruction=None, standard="기준", applicable=True):
    return Criterion(
        item_id=item_id, item_name="항목", risk=3.0, variant=variant,
        eval_type="스크립트", standard=standard, method="방법",
        applicable=applicable, judgment_method=judgment_method,
        summary_instruction=summary_instruction)


def test_needs_llm_true_for_llm_method():
    criteria = {("X-001", "AWS"): _crit(judgment_method="llm")}
    assert main_mod._needs_llm(criteria, "AWS") is True


def test_needs_llm_true_for_llm_det_method():
    criteria = {("X-001", "AWS"): _crit(judgment_method="llm_det")}
    assert main_mod._needs_llm(criteria, "AWS") is True


def test_needs_llm_false_for_pure_det_and_fw_policy_only():
    """iss 프로파일처럼 전 항목이 det/fw_policy(순결정론)뿐이면 LLM 불필요."""
    criteria = {
        ("X-001", "fw"): _crit(variant="fw", judgment_method="det"),
        ("X-002", "fw"): _crit(variant="fw", judgment_method="fw_policy"),
    }
    assert main_mod._needs_llm(criteria, "fw") is False


def test_needs_llm_false_for_interview_holdonly_without_summary():
    criteria = {("X-001", "AWS"): _crit(
        judgment_method="interview_holdonly", summary_instruction=None)}
    assert main_mod._needs_llm(criteria, "AWS") is False


def test_needs_llm_true_for_interview_with_summary_instruction():
    """method 이름이 interview/interview_holdonly여도 실제
    summary_instruction이 채워져 있으면(실행 시 LLM 요약 호출) LLM 필요로 본다."""
    criteria = {("X-001", "AWS"): _crit(
        judgment_method="interview", summary_instruction="요약하라")}
    assert main_mod._needs_llm(criteria, "AWS") is True


def test_needs_llm_ignores_non_judgeable_and_other_variant():
    criteria = {
        # 다른 variant는 무시
        ("X-001", "Azure"): _crit(variant="Azure", judgment_method="llm"),
        # applicable=False(비대상)는 is_judgeable=False → 무시
        ("X-002", "AWS"): _crit(variant="AWS", judgment_method="llm",
                                applicable=False),
        ("X-003", "AWS"): _crit(variant="AWS", judgment_method="det"),
    }
    assert main_mod._needs_llm(criteria, "AWS") is False


def test_needs_llm_false_on_empty_criteria():
    assert main_mod._needs_llm({}, "AWS") is False


# --------------------------------------------------------------------------
# 스마트 런처: --criteria 자동탐색 (_discover_criteria_path)
# --------------------------------------------------------------------------

def test_discover_criteria_path_no_match_raises_report_error(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    # 패키지 루트 폴백도 실데이터를 못 찾도록 존재하지 않는 경로로 돌린다.
    monkeypatch.setattr(main_mod, "_PKG_ROOT", str(tmp_path / "no_such_pkg_root"))
    with pytest.raises(ReportError):
        main_mod._discover_criteria_path()


def test_discover_criteria_path_single_match(tmp_path, monkeypatch):
    ref_dir = tmp_path / "ref"
    ref_dir.mkdir()
    f = ref_dir / "전자금융기반시설 보안 취약점 평가기준(제2026-1호) 평가자용.xlsx"
    f.write_text("x", encoding="utf-8")
    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(main_mod, "_PKG_ROOT", str(tmp_path / "no_such_pkg_root"))

    found = main_mod._discover_criteria_path()
    assert os.path.realpath(found) == os.path.realpath(str(f))


def test_discover_criteria_path_picks_latest_version_when_multiple(
        tmp_path, monkeypatch):
    ref_dir = tmp_path / "ref"
    ref_dir.mkdir()
    old = ref_dir / "평가기준(제2025-3호).xlsx"
    new = ref_dir / "평가기준(제2026-1호).xlsx"
    old.write_text("x", encoding="utf-8")
    new.write_text("x", encoding="utf-8")
    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(main_mod, "_PKG_ROOT", str(tmp_path / "no_such_pkg_root"))

    found = main_mod._discover_criteria_path()
    assert os.path.realpath(found) == os.path.realpath(str(new))


# --------------------------------------------------------------------------
# 스마트 런처: --profile 자동추정 실패 시 종료 (_resolve_profile)
# --------------------------------------------------------------------------

def test_resolve_profile_explicit_wins_over_guess():
    # 명시 지정이 있으면 파일명이 무엇이든 그대로 사용(기존 동작 불변).
    assert main_mod._resolve_profile("random.xml", "cloud") == "cloud"


def test_resolve_profile_guesses_from_marker():
    assert main_mod._resolve_profile("aws_report_x.xml", None) == "cloud"


def test_resolve_profile_ambiguous_raises_report_error():
    with pytest.raises(ReportError, match="확정할 수 없습니다"):
        main_mod._resolve_profile("unknown_result.json", None)


def test_resolve_profile_unrecognized_raises_report_error():
    with pytest.raises(ReportError, match="추정할 수 없습니다"):
        main_mod._resolve_profile("notes.txt", None)


# --------------------------------------------------------------------------
# 스마트 런처: main() 하위호환 — --profile/--criteria 명시 시 자동추정 우회
# --------------------------------------------------------------------------

def test_main_cli_profile_omitted_is_autoguessed_from_filename(
        tmp_path, monkeypatch):
    """--profile을 생략해도 파일명(aws_report_*)으로 cloud가 자동추정되어
    기존처럼 동작해야 한다(하위호환: 명시 호출은 그대로 동작)."""
    rep_dir, crit_dir, out_dir = (tmp_path / d for d in ("rep", "crit", "out"))
    for d in (rep_dir, crit_dir, out_dir):
        os.makedirs(str(d), exist_ok=True)
    report = _copy_fixture_xml(rep_dir)
    criteria = os.path.join(str(crit_dir), "criteria.xlsx")
    _write_synthetic_criteria(criteria)

    monkeypatch.setattr(main_mod, "OllamaClient",
                        lambda *a, **k: StubClient())

    main(["--report", report, "--criteria", criteria,
          "--out-dir", str(out_dir), "--model", "stub"])

    base = os.path.splitext(os.path.basename(report))[0]
    assert os.path.exists(os.path.join(str(out_dir), f"result_{base}.json"))


def test_main_positional_report_argument_works(tmp_path, monkeypatch):
    """위치인자로 --report를 대체할 수 있어야 한다(python3 -m judge_tool <파일>)."""
    rep_dir, crit_dir, out_dir = (tmp_path / d for d in ("rep", "crit", "out"))
    for d in (rep_dir, crit_dir, out_dir):
        os.makedirs(str(d), exist_ok=True)
    report = _copy_fixture_xml(rep_dir)
    criteria = os.path.join(str(crit_dir), "criteria.xlsx")
    _write_synthetic_criteria(criteria)

    monkeypatch.setattr(main_mod, "OllamaClient",
                        lambda *a, **k: StubClient())

    main([report, "--criteria", criteria, "--out-dir", str(out_dir),
          "--model", "stub"])

    base = os.path.splitext(os.path.basename(report))[0]
    assert os.path.exists(os.path.join(str(out_dir), f"result_{base}.json"))


def test_main_no_report_at_all_exits_cleanly():
    with pytest.raises(SystemExit):
        main([])


def test_main_profile_ambiguous_exits_with_candidates(tmp_path, capsys):
    """마커/확장자로 프로파일을 확정할 수 없는 입력은 후보를 안내하고 종료한다
    (오판정보다 명시 요구가 안전 — cloud 무조건 폴백 제거 확인)."""
    report = os.path.join(str(tmp_path), "unknown_result.json")
    with open(report, "w", encoding="utf-8") as fh:
        fh.write("{}")
    criteria = os.path.join(str(tmp_path), "criteria.xlsx")
    _write_synthetic_criteria(criteria)

    with pytest.raises(SystemExit):
        main(["--report", report, "--criteria", criteria])

    err = capsys.readouterr().err
    assert "확정할 수 없습니다" in err


# --------------------------------------------------------------------------
# 스마트 런처: 배치 모드 후보파일 나열 (_list_candidate_reports)
# --------------------------------------------------------------------------

def test_list_candidate_reports_filters_criteria_and_prior_outputs(tmp_path):
    (tmp_path / "aws_report_a.xml").write_text("x", encoding="utf-8")
    (tmp_path / "mysql_result.json").write_text("{}", encoding="utf-8")
    # DB 결과 표준 포맷 .txt(마커 있는 실결과)는 후보에 포함되어야 한다.
    (tmp_path / "oracle_result.txt").write_text("x", encoding="utf-8")
    (tmp_path / "평가기준(제2026-1호).xlsx").write_text("x", encoding="utf-8")
    (tmp_path / "result_aws_report_a.json").write_text("{}", encoding="utf-8")
    (tmp_path / ".hidden.xml").write_text("x", encoding="utf-8")

    found = {os.path.basename(p)
            for p in main_mod._list_candidate_reports(str(tmp_path))}
    # 기준 xlsx / result_ 접두 산출물 / 숨김파일은 제외, .txt 결과는 포함.
    assert found == {"aws_report_a.xml", "mysql_result.json", "oracle_result.txt"}


def test_batch_explicit_outdir_inside_scan_rejected_before_mkdir(tmp_path):
    """M-2: 명시적 --out-dir이 스캔 폴더 내부면 디렉터리 생성 전에 거부한다."""
    import types
    scan = tmp_path / "scan"
    scan.mkdir()
    (scan / "aws_report_a.xml").write_text("x", encoding="utf-8")
    (tmp_path / "crit.xlsx").write_text("x", encoding="utf-8")
    bad_out = scan / "out"  # 스캔 폴더 하위 → 실데이터 오염 위험 → 거부 대상
    ns = types.SimpleNamespace(
        report=str(scan), out_dir=str(bad_out), criteria=str(tmp_path / "crit.xlsx"),
        profile=None, ollama_url="http://x:11434", model="qwen3-coder:30b",
        skip_preflight=True, hashcat_path=None, hashcat_wordlist=None,
        hashcat_rules=None, hashcat_timeout=600)
    with pytest.raises(SystemExit):
        main_mod._run_batch(ns)
    # 가드가 makedirs보다 먼저 걸리므로 out 디렉터리가 만들어지지 않아야 한다.
    assert not bad_out.exists()


def test_extract_criteria_version_other_version():
    assert _extract_criteria_version("기준(제2027-3호).xlsx") == "제2027-3호"


def test_extract_criteria_version_multi_match_takes_first():
    """다중 매칭 시 첫 매칭을 채택한다(현 동작 고정)."""
    p = "기준(제2026-1호)_(제2027-2호).xlsx"
    assert _extract_criteria_version(p) == "제2026-1호"


def test_run_empty_input(tmp_path):
    """증거가 0건인 보고서 → 판정대상 전 항목이 '증거 미수집' 자동보류로
    출력된다(조용한 누락 금지). 출력 정상 생성, 예외 없음."""
    report = os.path.join(str(tmp_path), "aws_report_empty.xml")
    with open(report, "w", encoding="utf-8") as fh:
        fh.write('<?xml version="1.0" encoding="UTF-8"?>\n'
                 "<AuditReport><CheckList></CheckList></AuditReport>")
    criteria = os.path.join(str(tmp_path), "criteria.xlsx")
    _write_synthetic_criteria(criteria)
    json_out = os.path.join(str(tmp_path), "result.json")
    xlsx_out = os.path.join(str(tmp_path), "result.xlsx")

    cov = run(report_path=report, criteria_path=criteria, profile_key="cloud",
              client=StubClient(), json_out=json_out, xlsx_out=xlsx_out,
              model_name="stub")

    assert os.path.exists(json_out) and os.path.exists(xlsx_out)
    # 판정대상 3건(PISM-001/037/007) 모두 증거 미수집 자동보류로 출력
    assert cov["judged"] == 3
    assert cov["missing"] == []
    data = json.load(open(json_out, encoding="utf-8"))
    assert len(data["judgments"]) == 3
    for j in data["judgments"]:
        assert j["verdict"] == "판단보류"
        assert "증거 미수집" in j["rationale"]
        assert j["needs_review"] is True
