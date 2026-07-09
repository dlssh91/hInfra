"""fw_policy_xlsx 파서 단위 테스트 (작업②).

합성 openpyxl xlsx 픽스처로 실데이터 없이 실행한다.
"""
import csv
import json

import openpyxl
import pytest

from judge_tool.errors import ReportError
from judge_tool.parsers.fw_policy_xlsx import (
    detect_variant,
    parse,
    _ISS_FW_IDS,
    _detect_format_from_rows,
    _parse_id70,
    _parse_krfw,
    _parse_paloalto,
    _parse_secui,
)


# ─ detect_variant ─────────────────────────────────────────────────────────────

def test_detect_variant_always_fw():
    assert detect_variant("any_path.xlsx") == "fw"
    assert detect_variant("") == "fw"


# ─ 합성 xlsx 헬퍼 ──────────────────────────────────────────────────────────────

def _write_xlsx(path, sheet_name, rows):
    """rows: list of list; 각 원소가 셀값."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name
    for r_idx, row in enumerate(rows, start=1):
        for c_idx, val in enumerate(row, start=1):
            ws.cell(r_idx, c_idx, val)
    wb.save(str(path))


def _secui_rows():
    """최소 SECUI 시트 rows (헤더 + 2개 정책)."""
    return [
        # row1: 범례 (font color 키워드)
        ["Font Color:", "색", None, None, None],
        # row2: 빈행
        [None]*5,
        # row3: 주요 헤더 행
        ["Seq", "Enable", "Two-way", "ID", "Action", None, None, None, "IP", None, None,
         None, None, None, None, "IP", None, None, None, None, "Protocol", "Port"],
        # row4: 데이터 행1 — any-any allow
        [1, "Y", "N", "R001", "allow", None, None, None, "any", None, None,
         None, None, None, None, "any", None, None, None, None, "tcp", "80"],
        # row5: 데이터 행2 — deny
        [2, "Y", "N", "R002", "deny", None, None, None, "10.0.0.1", None, None,
         None, None, None, None, "192.168.1.1", None, None, None, None, "tcp", "443"],
    ]


def _id70_rows():
    """최소 ID70 시트 rows (헤더 + 2개 정책)."""
    return [
        ["PRIORITY", "ENABLED", "SRC ADDR", "DST ADDR", "SVC SPEC", "ACTION",
         "DAILY HIT COUNT"],
        ["P001", "YES", "any", "10.0.0.0/8", "tcp 1-65535 22-22", "allow", 0],
        ["P002", "YES", "192.168.1.0/24", "any", "tcp 1-65535 443-443", "deny", 100],
    ]


def _paloalto_rows():
    """최소 Palo Alto 시트 rows (헤더 + 2개 정책)."""
    return [
        ["Name", "Source", "Destination", "Service", "Action", "Hit Count", "Application"],
        ["rule1", "any", "10.0.0.0/8", "22", "allow", 0, "ssh"],
        ["rule2", "192.168.1.0/24", "any", "443", "deny", 50, "ssl"],
    ]


def _krfw_rows():
    """최소 krfw(한글 13열) 시트 rows (헤더 + 2개 정책). 더미IP만 사용."""
    header = ["룰 NUM", "출발지", "목적지", "서비스", "Protocol", "inbound",
              "시간", "정책", "로그", "session-limit", "tcp", "활성화", "설명"]
    return [
        header,
        ["1\n", "10.0.0.1", "192.168.1.1", "443", "tcp", None, None,
         "허용", "Enable", None, None, "Enable", "테스트 규칙1"],
        ["2", "any", "any", None, "icmp", None, None,
         "차단", "Enable", None, None, "off", "테스트 규칙2"],
    ]


# ─ _detect_format_from_rows ───────────────────────────────────────────────────

def test_detect_format_secui():
    rows = [(v,) if not isinstance(v, (list, tuple)) else tuple(v)
            for v in _secui_rows()]
    rows_as_tuples = [tuple(r) if isinstance(r, list) else r for r in _secui_rows()]
    fmt = _detect_format_from_rows(rows_as_tuples)
    assert fmt == "secui"


def test_detect_format_id70():
    rows_as_tuples = [tuple(r) for r in _id70_rows()]
    fmt = _detect_format_from_rows(rows_as_tuples)
    assert fmt == "id70"


def test_detect_format_paloalto():
    rows_as_tuples = [tuple(r) for r in _paloalto_rows()]
    fmt = _detect_format_from_rows(rows_as_tuples)
    assert fmt == "paloalto"


def test_detect_format_all_empty():
    fmt = _detect_format_from_rows([(None, None, None)] * 10)
    assert fmt == "unknown"


def test_detect_format_krfw():
    rows_as_tuples = [tuple(r) for r in _krfw_rows()]
    fmt = _detect_format_from_rows(rows_as_tuples)
    assert fmt == "krfw"


# ─ _parse_id70 단위 ───────────────────────────────────────────────────────────

def test_parse_id70_policies():
    rows = [tuple(r) for r in _id70_rows()]
    policies = _parse_id70(rows, "test")
    assert len(policies) == 2
    p = policies[0]
    assert p.rule_id == "P001"
    assert p.action == "allow"
    assert p.src_ips == ["any"]
    assert p.dst_ips == ["10.0.0.0/8"]
    assert p.hit_count == 0


def test_parse_id70_svc_spec_parsed():
    rows = [tuple(r) for r in _id70_rows()]
    policies = _parse_id70(rows, "test")
    p = policies[0]  # "tcp 1-65535 22-22"
    assert "tcp" in p.protocols
    assert "22-22" in p.dst_ports


def test_parse_id70_hit_count():
    rows = [tuple(r) for r in _id70_rows()]
    policies = _parse_id70(rows, "test")
    assert policies[0].hit_count == 0
    assert policies[1].hit_count == 100


def test_parse_id70_empty_rows():
    assert _parse_id70([], "test") == []


def test_parse_id70_no_priority_col():
    rows = [("Name", "Action", "Source"), ("rule1", "allow", "any")]
    assert _parse_id70([tuple(r) for r in rows], "test") == []


# ─ _parse_secui 단위 ───────────────────────────────────────────────────────────

def test_parse_secui_policies():
    rows = [tuple(r) for r in _secui_rows()]
    policies = _parse_secui(rows, "test")
    assert len(policies) == 2
    assert policies[0].action == "allow"
    assert policies[1].action == "deny"


def test_parse_secui_seq():
    rows = [tuple(r) for r in _secui_rows()]
    policies = _parse_secui(rows, "test")
    assert policies[0].seq == 1
    assert policies[1].seq == 2


def test_parse_secui_ips_extracted():
    rows = [tuple(r) for r in _secui_rows()]
    policies = _parse_secui(rows, "test")
    # any-any row
    p0 = policies[0]
    assert "any" in p0.src_ips or len(p0.src_ips) == 0 or "any" in p0.src_ips


def test_parse_secui_src_ports_not_collected():
    """(B'-2 증거조작 제거) SECUI는 출발지 포트 컬럼이 없어 src_ports는
    항상 []（["any"] 하드코딩 금지 — capability=False 게이트가 담당)."""
    rows = [tuple(r) for r in _secui_rows()]
    policies = _parse_secui(rows, "test")
    for p in policies:
        assert p.src_ports == []


# ─ _parse_paloalto 단위 ───────────────────────────────────────────────────────

def test_parse_paloalto_policies():
    rows = [tuple(r) for r in _paloalto_rows()]
    policies = _parse_paloalto(rows, "test")
    assert len(policies) == 2


def test_parse_paloalto_action_and_ips():
    rows = [tuple(r) for r in _paloalto_rows()]
    policies = _parse_paloalto(rows, "test")
    assert policies[0].action == "allow"
    assert policies[0].src_ips == ["any"]
    assert policies[0].dst_ips == ["10.0.0.0/8"]
    assert policies[0].hit_count == 0


def test_parse_paloalto_empty_rows():
    assert _parse_paloalto([], "test") == []


def test_parse_paloalto_skip_empty_action():
    rows = [
        ("Source", "Action", "Hit Count"),
        ("any", None, 0),   # action 없음 → skip
        ("any", "allow", 0),
    ]
    policies = _parse_paloalto([tuple(r) for r in rows], "test")
    assert len(policies) == 1


# ─ _parse_krfw 단위 ────────────────────────────────────────────────────────────

def test_parse_krfw_policies_count():
    rows = [tuple(r) for r in _krfw_rows()]
    policies = _parse_krfw(rows, "test")
    assert len(policies) == 2


def test_parse_krfw_field_mapping():
    rows = [tuple(r) for r in _krfw_rows()]
    policies = _parse_krfw(rows, "test")
    p0 = policies[0]
    assert p0.seq == 1
    assert p0.src_ips == ["10.0.0.1"]
    assert p0.dst_ips == ["192.168.1.1"]
    assert p0.dst_ports == ["443"]
    assert p0.protocols == ["tcp"]
    assert p0.description == "테스트 규칙1"


def test_parse_krfw_src_ports_not_collected():
    """(B'-2 증거조작 제거) krfw는 출발지 포트 컬럼이 없어 src_ports는
    항상 []（["any"] 하드코딩 금지 — capability=False 게이트가 담당)."""
    rows = [tuple(r) for r in _krfw_rows()]
    policies = _parse_krfw(rows, "test")
    for p in policies:
        assert p.src_ports == []


def test_parse_krfw_action_mapping_allow_deny():
    rows = [tuple(r) for r in _krfw_rows()]
    policies = _parse_krfw(rows, "test")
    assert policies[0].action == "allow"  # "허용" → allow
    assert policies[1].action == "deny"   # "차단" → deny


def test_parse_krfw_enabled_mapping():
    rows = [tuple(r) for r in _krfw_rows()]
    policies = _parse_krfw(rows, "test")
    assert policies[0].enabled is True   # "Enable" → True
    assert policies[1].enabled is False  # "off" → False


def test_parse_krfw_action_case_and_space_insensitive():
    header = _krfw_rows()[0]
    rows = [
        header,
        ["1", "any", "any", None, "tcp", None, None, "  PERMIT  ", "Enable",
         None, None, "Y", None],
        ["2", "any", "any", None, "tcp", None, None, "DENY", "Enable",
         None, None, "N", None],
    ]
    policies = _parse_krfw([tuple(r) for r in rows], "test")
    assert policies[0].action == "allow"
    assert policies[0].enabled is True
    assert policies[1].action == "deny"
    assert policies[1].enabled is False


def test_parse_krfw_unrecognized_action_tracked_in_stats():
    header = _krfw_rows()[0]
    rows = [
        header,
        ["1", "any", "any", None, "tcp", None, None, "이상한값", "Enable",
         None, None, "Enable", None],
        ["2", "any", "any", None, "tcp", None, None, "허용", "Enable",
         None, None, "Enable", None],
    ]
    stats = {}
    policies = _parse_krfw([tuple(r) for r in rows], "test", stats)
    assert len(policies) == 2
    assert stats.get("unrecognized_action_count") == 1
    # 미인식 값은 allow로 세지 않는다(원문 보존, allow로 취급 안 함)
    assert policies[0].action == "이상한값"


def test_parse_krfw_enabled_unrecognized_defaults_true():
    header = _krfw_rows()[0]
    rows = [
        header,
        ["1", "any", "any", None, "tcp", None, None, "허용", "Enable",
         None, None, "미확인값", None],
    ]
    policies = _parse_krfw([tuple(r) for r in rows], "test")
    assert policies[0].enabled is True


def test_parse_krfw_empty_rows():
    assert _parse_krfw([], "test") == []


def test_parse_krfw_wrong_headers_returns_empty():
    rows = [("Name", "Action", "Source"), ("rule1", "allow", "any")]
    assert _parse_krfw([tuple(r) for r in rows], "test") == []


def test_parse_krfw_header_not_first_row():
    """(C-1②) 제목행 등으로 헤더가 rows[0]이 아니어도 상단 수 행에서 탐색."""
    title = ["방화벽 정책 현황"] + [None] * 12
    rows = [title] + _krfw_rows()
    policies = _parse_krfw([tuple(r) for r in rows], "test")
    assert len(policies) == 2
    assert policies[0].src_ips == ["10.0.0.1"]
    assert policies[0].action == "allow"


def test_parse_krfw_header_beyond_scan_window_returns_empty():
    """헤더가 탐색 창(상단 10행)을 벗어나면 파싱 0건(→ parse()가 unknown 강등)."""
    filler = [["필러"] + [None] * 12 for _ in range(12)]
    rows = filler + _krfw_rows()
    assert _parse_krfw([tuple(r) for r in rows], "test") == []


# ─ parse() — 통합 ──────────────────────────────────────────────────────────────

def test_parse_returns_12_tuples_secui(tmp_path):
    p = str(tmp_path / "secui.xlsx")
    _write_xlsx(p, "FW_Policy", _secui_rows())
    result = parse(p)
    assert len(result) == 12
    ids = [r[0] for r in result]
    assert ids == _ISS_FW_IDS


def test_parse_resources_always_empty(tmp_path):
    p = str(tmp_path / "secui2.xlsx")
    _write_xlsx(p, "FW_Policy", _secui_rows())
    result = parse(p)
    for iss_id, resources, ctx in result:
        assert resources == []


def test_parse_context_format_secui(tmp_path):
    p = str(tmp_path / "secui3.xlsx")
    _write_xlsx(p, "FW_Policy", _secui_rows())
    result = parse(p)
    ctx = result[0][2]
    assert ctx.startswith("FW_FORMAT:secui")
    assert "FW_SHEETS:" in ctx
    assert "FW_POLICY_COUNT:" in ctx
    assert "FW_POLICIES_JSON:" in ctx


def test_parse_context_json_valid(tmp_path):
    p = str(tmp_path / "id70.xlsx")
    _write_xlsx(p, "FW_Policy", _id70_rows())
    result = parse(p)
    ctx = result[0][2]
    json_part = ctx.split("FW_POLICIES_JSON:", 1)[1]
    policies = json.loads(json_part)
    assert isinstance(policies, list)
    assert len(policies) == 2


def test_parse_id70_file(tmp_path):
    p = str(tmp_path / "id70b.xlsx")
    _write_xlsx(p, "FW_Policy", _id70_rows())
    result = parse(p)
    ctx = result[0][2]
    assert "FW_FORMAT:id70" in ctx


def test_parse_paloalto_file(tmp_path):
    p = str(tmp_path / "paloalto.xlsx")
    _write_xlsx(p, "FW_Policy", _paloalto_rows())
    result = parse(p)
    ctx = result[0][2]
    assert "FW_FORMAT:paloalto" in ctx


def test_parse_skip_sheets_meta(tmp_path):
    """'점검대상' 시트는 스킵되고 정책 시트만 파싱된다."""
    wb = openpyxl.Workbook()
    # 첫 시트를 '점검대상'으로
    ws1 = wb.active
    ws1.title = "점검대상"
    for r, row in enumerate(_id70_rows(), start=1):
        for c, val in enumerate(row, start=1):
            ws1.cell(r, c, val)
    # 두 번째 시트에 실제 정책
    ws2 = wb.create_sheet("방화벽정책")
    for r, row in enumerate(_id70_rows(), start=1):
        for c, val in enumerate(row, start=1):
            ws2.cell(r, c, val)
    p = str(tmp_path / "multi.xlsx")
    wb.save(p)
    result = parse(p)
    ctx = result[0][2]
    # 방화벽정책 시트만 포함
    assert "방화벽정책" in ctx
    assert "점검대상" not in ctx.split("FW_SHEETS:")[1].split("\n")[0]


def test_parse_no_valid_sheet_downgrades_to_unknown(tmp_path):
    """빈 시트만 있으면 (B'-1) 크래시 대신 unknown+판단보류로 강등된다.

    이전 동작(ReportError raise)은 P13/P14/P24/P25 실증에서 크래시를
    유발했음이 확인되어(설계문서 §Phase A) 방어적 강등으로 교체됨.
    파일 자체를 열 수 없는 경우(별도 테스트)는 여전히 ReportError.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "빈시트"
    # 아무 내용 없음
    p = str(tmp_path / "empty.xlsx")
    wb.save(p)
    result = parse(p)
    assert len(result) == 12
    ctx = result[0][2]
    assert "FW_FORMAT:unknown" in ctx
    assert "FW_POLICY_COUNT:0" in ctx


def test_parse_unrecognized_format_no_crash(tmp_path):
    """헤더가 4개 포맷 중 어느 것과도 매칭되지 않으면 크래시 없이 unknown emit."""
    rows = [
        ["PolicyName", "Status", "Port", "Owner"],
        ["r1", "active", "80", "teamA"],
    ]
    p = str(tmp_path / "unknown_fmt.xlsx")
    _write_xlsx(p, "정책", rows)
    result = parse(p)
    assert len(result) == 12
    ctx = result[0][2]
    assert "FW_FORMAT:unknown" in ctx
    assert "FW_POLICY_COUNT:0" in ctx


def test_parse_sniffed_format_but_zero_policies_downgrades_to_unknown(tmp_path):
    """(C-1①) sniff는 krfw로 인식했지만 파싱 결과 정책 0건 → unknown 강등.

    거짓양호 봉쇄: fmt가 유지되면 capability 있는 항목이 '정책 0건=양호'로
    나가므로, 전 시트 합산 0건이면 unknown으로 강등해 전 항목 판단보류.
    """
    # krfw 헤더만 있고 데이터 행 없음 → sniff=krfw, 정책 0건
    rows = [_krfw_rows()[0]]
    p = str(tmp_path / "krfw_headeronly.xlsx")
    _write_xlsx(p, "Detail", rows)
    result = parse(p)
    assert len(result) == 12
    ctx = result[0][2]
    assert "FW_FORMAT:unknown" in ctx
    assert "FW_POLICY_COUNT:0" in ctx


def test_parse_title_row_krfw_file_parses_normally(tmp_path):
    """(C-1②) 1행 제목 + 2행 krfw 헤더 + 3행 정책 → 정상 파싱(0건 강등 아님)."""
    title = ["방화벽 정책 현황"] + [None] * 12
    rows = [title] + _krfw_rows()
    p = str(tmp_path / "krfw_title.xlsx")
    _write_xlsx(p, "Detail", rows)
    result = parse(p)
    ctx = result[0][2]
    assert "FW_FORMAT:krfw" in ctx
    assert "FW_POLICY_COUNT:2" in ctx


def test_parse_krfw_file_full_pipeline(tmp_path):
    p = str(tmp_path / "krfw.xlsx")
    _write_xlsx(p, "Detail", _krfw_rows())
    result = parse(p)
    ctx = result[0][2]
    assert "FW_FORMAT:krfw" in ctx
    assert "FW_POLICY_COUNT:2" in ctx
    json_part = ctx.split("FW_POLICIES_JSON:", 1)[1]
    policies = json.loads(json_part)
    assert len(policies) == 2


# ─ parse_stats 계측 ────────────────────────────────────────────────────────────

def test_parse_context_includes_parse_stats(tmp_path):
    p = str(tmp_path / "stats.xlsx")
    _write_xlsx(p, "Detail", _krfw_rows())
    result = parse(p)
    ctx = result[0][2]
    assert "FW_PARSE_STATS_JSON:" in ctx
    stats_part = ctx.split("FW_PARSE_STATS_JSON:", 1)[1].split("\nFW_POLICIES_JSON:", 1)[0]
    stats = json.loads(stats_part)
    for key in ("rows_total", "rows_parsed", "rows_dropped", "unrecognized_action_count"):
        assert key in stats


def test_parse_stats_unrecognized_action_count_propagated(tmp_path):
    header = _krfw_rows()[0]
    rows = [
        header,
        ["1", "any", "any", None, "tcp", None, None, "이상한값", "Enable",
         None, None, "Enable", None],
    ]
    p = str(tmp_path / "stats_unrecog.xlsx")
    _write_xlsx(p, "Detail", rows)
    result = parse(p)
    ctx = result[0][2]
    stats_part = ctx.split("FW_PARSE_STATS_JSON:", 1)[1].split("\nFW_POLICIES_JSON:", 1)[0]
    stats = json.loads(stats_part)
    assert stats["unrecognized_action_count"] == 1


# ─ CSV 분기 ────────────────────────────────────────────────────────────────────

def test_parse_csv_krfw_loads(tmp_path):
    p = tmp_path / "P99_정책.csv"
    with open(p, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        for row in _krfw_rows():
            writer.writerow(["" if v is None else v for v in row])
    result = parse(str(p))
    assert len(result) == 12
    ctx = result[0][2]
    assert "FW_FORMAT:krfw" in ctx
    assert "FW_POLICY_COUNT:2" in ctx


def test_parse_csv_cp949_fallback(tmp_path):
    p = tmp_path / "cp949.csv"
    lines = [",".join(str(v) if v is not None else "" for v in row)
             for row in _krfw_rows()]
    content = "\n".join(lines)
    with open(p, "w", encoding="cp949", newline="") as f:
        f.write(content)
    result = parse(str(p))
    assert len(result) == 12
    ctx = result[0][2]
    assert "FW_FORMAT:krfw" in ctx


def test_parse_csv_unknown_format_no_crash(tmp_path):
    p = tmp_path / "P15_정책.csv"
    with open(p, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(["?", "??", "???"])
        writer.writerow(["1", "a", "b"])
    result = parse(str(p))
    assert len(result) == 12
    ctx = result[0][2]
    assert "FW_FORMAT:unknown" in ctx


def test_parse_invalid_csv_path_raises():
    with pytest.raises(ReportError):
        parse("/nonexistent/path/file.csv")


def test_parse_invalid_path_raises():
    with pytest.raises(ReportError):
        parse("/nonexistent/path/file.xlsx")


def test_parse_all_iss_ids_present(tmp_path):
    p = str(tmp_path / "all_ids.xlsx")
    _write_xlsx(p, "FW_Policy", _paloalto_rows())
    result = parse(p)
    returned_ids = {r[0] for r in result}
    assert returned_ids == set(_ISS_FW_IDS)


def test_parse_shared_context_across_all_items(tmp_path):
    """12개 3-튜플이 동일한 context 문자열을 공유한다."""
    p = str(tmp_path / "shared.xlsx")
    _write_xlsx(p, "FW_Policy", _id70_rows())
    result = parse(p)
    contexts = [ctx for _, _, ctx in result]
    assert all(c == contexts[0] for c in contexts)


# ─ resolve_policies 배선 (B′-3a) — parse() 말미에서 1회 호출 ─────────────────

def _find_policy_by_seq(policies, seq):
    return next(pol for pol in policies if pol.get("seq") == seq)


def test_parse_krfw_named_dst_object_moves_to_unresolved(tmp_path):
    """krfw 목적지 셀이 named 객체(그룹명)면 dst_ips가 아니라 unresolved_dst로."""
    header = _krfw_rows()[0]
    rows = [
        header,
        ["1", "10.0.0.1", "WEB_SERVERS_GRP", "443", "tcp", None, None,
         "허용", "Enable", None, None, "Enable", "테스트"],
    ]
    p = str(tmp_path / "krfw_named.xlsx")
    _write_xlsx(p, "Detail", rows)
    result = parse(p)
    ctx = result[0][2]
    json_part = ctx.split("FW_POLICIES_JSON:", 1)[1]
    policies = json.loads(json_part)
    pol = _find_policy_by_seq(policies, 1)
    assert pol["dst_ips"] == []
    assert pol["unresolved_dst"] == ["WEB_SERVERS_GRP"]


def test_parse_krfw_ip_dst_stays_resolved(tmp_path):
    """krfw 목적지가 정상 IP면 unresolved_dst는 비어있다(회귀 방지)."""
    p = str(tmp_path / "krfw_ip.xlsx")
    _write_xlsx(p, "Detail", _krfw_rows())
    result = parse(p)
    ctx = result[0][2]
    json_part = ctx.split("FW_POLICIES_JSON:", 1)[1]
    policies = json.loads(json_part)
    for pol in policies:
        assert pol["unresolved_dst"] == []
        assert pol["unresolved_src"] == []


def test_parse_id70_named_svc_spec_moves_to_unresolved_svc(tmp_path):
    """ID70 SVC SPEC 단일 비숫자 토큰(named 서비스객체 후보) → unresolved_svc."""
    rows = [
        ["PRIORITY", "ENABLED", "SRC ADDR", "DST ADDR", "SVC SPEC", "ACTION",
         "DAILY HIT COUNT"],
        ["P001", "YES", "any", "10.0.0.0/8", "HTTPS_SVC_GRP", "allow", 5],
    ]
    p = str(tmp_path / "id70_named_svc.xlsx")
    _write_xlsx(p, "FW_Policy", rows)
    result = parse(p)
    ctx = result[0][2]
    json_part = ctx.split("FW_POLICIES_JSON:", 1)[1]
    policies = json.loads(json_part)
    assert len(policies) == 1
    assert policies[0]["protocols"] == []
    assert policies[0]["unresolved_svc"] == ["HTTPS_SVC_GRP"]
