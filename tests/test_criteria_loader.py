import openpyxl
import pytest

from judge_tool.criteria_loader import load_criteria
from judge_tool.profile import CLOUD, Profile, VariantSpec


def _write_noncanonical_criteria(path):
    """비정규형 item_id(예 'pism_001')가 든 합성 평가기준 xlsx 생성.

    로더가 profile.normalize_id 로 정규화해 'PISM-001' 키로 join 가능하게
    저장하는지 검증하기 위함.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = CLOUD.sheet_name
    aws = CLOUD.variants["AWS"]
    r = CLOUD.data_start_row
    ws.cell(r, CLOUD.id_col, "pism_001")        # 비정규형
    ws.cell(r, CLOUD.name_col, "통신구간 암호화")
    ws.cell(r, CLOUD.risk_col, 5)
    ws.cell(r, aws.eval_type_col, "스크립트")
    ws.cell(r, aws.method_col, "방법1")
    ws.cell(r, aws.standard_col, "양호 기준 텍스트")
    wb.save(path)


def test_noncanonical_item_id_is_normalized(tmp_path):
    path = str(tmp_path / "noncanon.xlsx")
    _write_noncanonical_criteria(path)
    crit = load_criteria(path, CLOUD)
    # 비정규형 'pism_001' 이 정규형 'PISM-001' 키로 join 가능해야 함
    assert ("PISM-001", "AWS") in crit
    assert ("pism_001", "AWS") not in crit
    c = crit[("PISM-001", "AWS")]
    # Criterion.item_id 도 정규형으로 저장(reconcile/coverage 일관성)
    assert c.item_id == "PISM-001"


def _write_full_mapping_criteria(path):
    """CLOUD 포맷의 모든 컬럼 인덱스를 변별 가능한 값으로 채운 합성 xlsx.

    AWS/Azure 의 eval_type/standard/method 컬럼을 서로 다른 텍스트로 두어
    load_criteria 의 컬럼 인덱스 매핑이 정확한지 결정적으로 검증한다.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = CLOUD.sheet_name
    aws = CLOUD.variants["AWS"]
    az = CLOUD.variants["Azure"]
    r = CLOUD.data_start_row
    ws.cell(r, CLOUD.id_col, "pism_042")        # 비정규형 → 정규화 검증
    ws.cell(r, CLOUD.name_col, "  접근통제  ")  # 공백 trim 검증
    ws.cell(r, CLOUD.risk_col, 3)
    # AWS 컬럼
    ws.cell(r, aws.eval_type_col, "스크립트")
    ws.cell(r, aws.standard_col, "AWS 판단기준")
    ws.cell(r, aws.method_col, "AWS 판단방법")
    # Azure 컬럼(AWS 와 구분되는 값)
    ws.cell(r, az.eval_type_col, "관리체계, 스크립트")
    ws.cell(r, az.standard_col, "Azure 판단기준")
    ws.cell(r, az.method_col, "Azure 판단방법")
    wb.save(path)


def test_full_column_mapping_is_exact(tmp_path):
    """각 variant 의 eval_type/standard/method 가 올바른 컬럼에서 읽혀야 한다."""
    path = str(tmp_path / "fullmap.xlsx")
    _write_full_mapping_criteria(path)
    crit = load_criteria(path, CLOUD)

    aws = crit[("PISM-042", "AWS")]
    assert aws.item_id == "PISM-042"          # 정규화
    assert aws.item_name == "접근통제"        # trim
    assert aws.risk == 3.0
    assert aws.variant == "AWS"
    assert aws.eval_type == "스크립트"
    assert aws.standard == "AWS 판단기준"
    assert aws.method == "AWS 판단방법"

    az = crit[("PISM-042", "Azure")]
    assert az.variant == "Azure"
    assert az.eval_type == "관리체계, 스크립트"
    assert az.standard == "Azure 판단기준"
    assert az.method == "Azure 판단방법"
    # 컬럼이 섞이지 않았는지(교차 오염 방지)
    assert az.standard != aws.standard
    assert az.method != aws.method
    assert az.is_mixed is True and aws.is_mixed is False


def test_blank_id_rows_are_skipped(tmp_path):
    """item_id 가 빈 행은 건너뛴다(data_start_row 이후 빈 행 안전)."""
    path = str(tmp_path / "blank.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = CLOUD.sheet_name
    aws = CLOUD.variants["AWS"]
    # data_start_row 는 비워 두고 그 다음 행에 데이터
    r = CLOUD.data_start_row + 1
    ws.cell(r, CLOUD.id_col, "pism_005")
    ws.cell(r, CLOUD.name_col, "로그관리")
    ws.cell(r, CLOUD.risk_col, 2)
    ws.cell(r, aws.eval_type_col, "스크립트")
    ws.cell(r, aws.standard_col, "기준")
    ws.cell(r, aws.method_col, "방법")
    wb.save(path)

    crit = load_criteria(path, CLOUD)
    assert ("PISM-005", "AWS") in crit
    # 빈 첫 데이터행으로 생긴 가짜 키가 없어야 함
    assert all(k[0] for k in crit)


def test_loads_known_item(criteria_xlsx_path):
    crit = load_criteria(criteria_xlsx_path, CLOUD)
    c = crit[("PISM-001", "AWS")]
    assert c.item_name.strip().startswith("통신구간")
    assert c.risk == 5.0
    assert "스크립트" in c.eval_type
    assert "양호" in c.standard          # 판단기준 텍스트 존재
    assert c.variant == "AWS"


def test_mixed_item_flagged(criteria_xlsx_path):
    crit = load_criteria(criteria_xlsx_path, CLOUD)
    assert crit[("PISM-045", "AWS")].is_mixed is True


def test_both_variants_present(criteria_xlsx_path):
    crit = load_criteria(criteria_xlsx_path, CLOUD)
    assert ("PISM-001", "AWS") in crit
    assert ("PISM-001", "Azure") in crit
    # 73개 항목 × 2 variant
    item_ids = {k[0] for k in crit}
    assert len(item_ids) == 73


# ── applies_when_standard 분기 단위 테스트 ────────────────────────────────────

def _make_applies_when_standard_profile(sheet="TestNet"):
    """applies_when_standard=True variant를 가진 합성 프로파일."""
    return Profile(
        key="testnet",
        sheet_name=sheet,
        header_row=4,
        data_start_row=5,
        id_col=2,
        name_col=7,
        risk_col=8,
        parser="network_xml",
        evidence_mode="raw",
        status_available=False,
        flag_vulnerable_for_review=True,
        empty_means_good=frozenset(),
        variants={
            "generic": VariantSpec(
                "generic", standard_col=5, method_col=6,
                applicability_col=None, applies_when_standard=True,
                filename_markers=()),
            # applicability_col 동시 설정 테스트용
            "with_col": VariantSpec(
                "with_col", standard_col=5, method_col=6,
                applicability_col=10, applies_when_standard=True,
                filename_markers=()),
        },
    )


def _write_applies_xlsx(tmp_path, standard_val, col10_val=""):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "TestNet"
    r = 5
    ws.cell(r, 2, "NET-001")
    ws.cell(r, 7, "테스트항목")
    ws.cell(r, 8, 5.0)
    ws.cell(r, 5, standard_val)
    ws.cell(r, 6, "방법")
    if col10_val:
        ws.cell(r, 10, col10_val)
    p = str(tmp_path / "applies.xlsx")
    wb.save(p)
    return p


def test_applies_when_standard_nonempty_applicable(tmp_path):
    """applies_when_standard=True, standard 비공백 → applicable=True."""
    p = _write_applies_xlsx(tmp_path, "* 양호 - 기준")
    profile = _make_applies_when_standard_profile()
    crit = load_criteria(p, profile)
    g = crit[("NET-001", "generic")]
    assert g.applicable is True
    assert g.is_judgeable is True


def test_applies_when_standard_empty_not_applicable(tmp_path):
    """applies_when_standard=True, standard 빈칸 → applicable=False."""
    p = _write_applies_xlsx(tmp_path, "")
    profile = _make_applies_when_standard_profile()
    crit = load_criteria(p, profile)
    g = crit[("NET-001", "generic")]
    assert g.applicable is False
    assert g.is_judgeable is False


def test_applicability_col_takes_precedence_over_applies_when_standard(tmp_path):
    """applicability_col 동시 설정 시 applicability_col('o') 이 우선 적용된다."""
    import os
    p = _write_applies_xlsx(tmp_path, "기준있음", col10_val="o")
    profile = _make_applies_when_standard_profile()
    crit = load_criteria(p, profile)
    # with_col: applicability_col=10, 'o' → applicable True
    wc = crit[("NET-001", "with_col")]
    assert wc.applicable is True

    # col10이 없으면 False (applicability_col 분기가 먼저라 applies_when_standard 무관)
    sub = tmp_path / "sub"
    os.makedirs(str(sub), exist_ok=True)
    p2 = _write_applies_xlsx(sub, "기준있음", col10_val="")
    crit2 = load_criteria(p2, profile)
    wc2 = crit2[("NET-001", "with_col")]
    assert wc2.applicable is False  # 'o' 없으면 False
