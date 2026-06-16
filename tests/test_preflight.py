"""tests/test_preflight.py — Pre-flight 인코딩 교정 + LLM 게이트 테스트.

설계 계약:
  - 인코딩 판별: utf-8 바이트+euc-kr 선언 → utf-8 선택.
  - 진짜 cp949/euc-kr 바이트 → cp949 선택.
  - utf-8-sig BOM → utf-8-sig 선택.
  - 깨진 바이트(셋 다 strict 실패) → replace 폴백.
  - 실데이터 컨테이너 XML(EUC-KR 선언 + 실제 UTF-8) → 교정 후 한글 정상.
  - LLM 게이트: fake OK → 진행. fake NG → PreflightError. LLM 예외 → 휴리스틱 폴백.
  - main.run의 --skip-preflight 동작 불변 확인.

실 Ollama 호출 없이 동작(fake 클라이언트만 사용).
"""
import json
import os
import re

import pytest

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CONTAINER_XML = os.path.join(
    PROJECT_ROOT,
    "collected", "container", "k8s_master",
    "fsec-control-plane-k8s_master-20260615.xml")


# ---------------------------------------------------------------------------
# 헬퍼 — fake LLM 클라이언트
# ---------------------------------------------------------------------------

class OKClient:
    """항상 OK를 반환하는 fake 게이트 클라이언트."""
    def chat(self, system, user):
        return json.dumps({"result": "OK", "reason": "테스트 양호"})


class NGClient:
    """항상 NG를 반환하는 fake 게이트 클라이언트."""
    def chat(self, system, user):
        return json.dumps({"result": "NG", "reason": "인코딩 깨짐 감지"})


class ExceptionClient:
    """항상 예외를 던지는 fake 클라이언트 (Ollama 미가동 시뮬레이션)."""
    def chat(self, system, user):
        raise ConnectionError("Ollama 연결 실패")


class JudgeStubClient:
    """main.run에서 LLM 판정용으로 사용되는 stub — 게이트에서도 OK를 반환."""
    def chat(self, system, user):
        # pre-flight 게이트 호출과 판정 호출을 모두 받아 처리
        # 게이트 프롬프트 여부 확인: system에 "result" 키가 언급되어 있으면 게이트
        if '"result"' in system:
            return json.dumps({"result": "OK", "reason": "stub 통과"})
        return json.dumps({
            "verdict": "양호", "confidence": 0.9,
            "rationale": "테스트 판정", "cited_evidence": ["x"]
        })


# ---------------------------------------------------------------------------
# Step 1: 인코딩 판별 테스트
# ---------------------------------------------------------------------------

class TestReadText:
    """preflight.read_text 의 인코딩 판별 동작."""

    def test_utf8_bytes_with_euckr_declaration_uses_utf8(self, tmp_path):
        """EUC-KR 선언이지만 실제 바이트가 UTF-8 → utf-8 선택 + 교정 기록."""
        from judge_tool.preflight import read_text

        # UTF-8 한글 포함 XML을 euc-kr로 선언
        xml_content = (
            '<?xml version="1.0" encoding="EUC-KR"?>\n'
            "<script><asset><hostname>테스트서버</hostname>"
            "<os>Linux</os></asset><results/></script>"
        )
        path = tmp_path / "test_fake_euckr.xml"
        path.write_bytes(xml_content.encode("utf-8"))

        text, meta = read_text(str(path))

        assert meta["used_encoding"] == "utf-8"
        assert meta["encoding_corrected"] is True
        assert meta["declared_encoding"] == "euc-kr"
        # 교정 후 한글이 정상
        assert "테스트서버" in text
        # XML 선언이 제거된 상태여야 ET.fromstring이 가능
        assert "<?xml" not in text

    def test_genuine_cp949_bytes_uses_cp949(self, tmp_path):
        """진짜 CP949(euc-kr 상위호환) 바이트 → cp949 선택."""
        from judge_tool.preflight import read_text

        xml_template = (
            '<?xml version="1.0" encoding="EUC-KR"?>\n'
            "<script><asset><hostname>서버이름</hostname>"
            "<os>Linux</os></asset><results/></script>"
        )
        path = tmp_path / "test_cp949.xml"
        path.write_bytes(xml_template.encode("cp949"))

        text, meta = read_text(str(path))

        assert meta["used_encoding"] == "cp949"
        # CP949는 strict UTF-8을 통과 못 하고 CP949를 선택
        assert meta["replace_fallback"] is False
        # 한글 정상
        assert "서버이름" in text

    def test_utf8_sig_bom_detected(self, tmp_path):
        """BOM(utf-8-sig) 파일 → utf-8-sig 선택."""
        from judge_tool.preflight import read_text

        xml_content = (
            '<?xml version="1.0" encoding="UTF-8"?>\n'
            "<script><results/></script>"
        )
        path = tmp_path / "test_bom.xml"
        path.write_bytes(b"\xef\xbb\xbf" + xml_content.encode("utf-8"))

        text, meta = read_text(str(path))

        assert meta["used_encoding"] == "utf-8-sig"
        assert meta["replace_fallback"] is False
        # 텍스트가 정상 파싱 가능
        assert "<script>" in text

    def test_broken_bytes_replace_fallback(self, tmp_path):
        """utf-8·cp949 모두 strict 실패하는 바이트 → replace 폴백."""
        from judge_tool.preflight import read_text

        # 어떤 인코딩으로도 디코딩 불가한 바이트 시퀀스 삽입
        invalid = (
            b'<?xml version="1.0" encoding="UTF-8"?>\n'
            b"<script><results/>"
            b"\x80\x81\x82\xff\xfe"  # 어떤 인코딩으로도 불법
            b"</script>"
        )
        path = tmp_path / "test_broken.xml"
        path.write_bytes(invalid)

        text, meta = read_text(str(path))

        assert meta["replace_fallback"] is True
        # replace 폴백은 크래시 없이 완료되어야 함
        assert isinstance(text, str)

    def test_real_container_xml_mojibake_corrected(self):
        """실데이터 컨테이너 XML(EUC-KR 선언 + 실제 UTF-8) 교정 확인.

        교정 전: euc-kr decode → mojibake(replacement char 다수).
        교정 후: utf-8 decode → 한글 정상.
        """
        if not os.path.exists(CONTAINER_XML):
            pytest.skip("컨테이너 실데이터 파일이 없어 스킵")

        from judge_tool.preflight import read_text

        text, meta = read_text(CONTAINER_XML)

        # 인코딩 교정이 일어났는지 확인
        assert meta["used_encoding"] == "utf-8"
        assert meta["encoding_corrected"] is True
        assert meta["declared_encoding"] == "euc-kr"

        # 교정 후 한글이 정상 (mojibake 없음)
        assert "불필요한" in text, "교정 후 한글 텍스트가 정상이어야 함"
        assert "클러스터" in text, "교정 후 한글 텍스트가 정상이어야 함"

        # replacement char(mojibake 마커)가 없어야 함
        replacement_count = text.count("■") + text.count("�")
        assert replacement_count == 0, \
            f"교정 후 replacement char가 없어야 함, found: {replacement_count}"

    def test_declared_encoding_matches_actual_no_correction(self, tmp_path):
        """선언 인코딩과 실제 인코딩이 일치 → encoding_corrected=False."""
        from judge_tool.preflight import read_text

        xml_content = (
            '<?xml version="1.0" encoding="UTF-8"?>\n'
            "<script><results/></script>"
        )
        path = tmp_path / "test_match.xml"
        path.write_bytes(xml_content.encode("utf-8"))

        text, meta = read_text(str(path))

        assert meta["encoding_corrected"] is False
        assert meta["used_encoding"] == "utf-8"

    def test_genuine_cp949_not_overtriggered_by_mojibake_guard(self):
        """점3 과트리거 방지: 진짜 한글 cp949 파일은 cp949로 선택되고 replace 폴백 없음.

        cp949 strict decode 후 U+FFFD 비율 0% → 5% 임계 통과 → cp949 정상 선택.
        실 한글 cp949 파일이 mojibake 가드로 잘못 차단되지 않아야 한다.
        """
        from judge_tool.preflight import _detect_encoding

        genuine_cp949_xml = (
            '<?xml version="1.0" encoding="EUC-KR"?>\n'
            "<script><asset><hostname>서버이름</hostname></asset></script>"
        )
        genuine_bytes = genuine_cp949_xml.encode("cp949")
        enc, fallback = _detect_encoding(genuine_bytes)

        assert enc == "cp949", f"진짜 cp949 파일은 cp949로 선택되어야 함, got {enc}"
        assert fallback is False, "진짜 cp949 파일은 replace 폴백이 아니어야 함"

    def test_sanitize_applied(self, tmp_path):
        """교정된 텍스트에 sanitize(불법 제어문자 제거·& escape)가 적용됨."""
        from judge_tool.preflight import read_text

        xml_content = (
            '<?xml version="1.0" encoding="UTF-8"?>\n'
            "<script><results><dump>"
            "<output><![CDATA[test \x1b[31m result & done]]></output>"
            "</dump></results></script>"
        )
        path = tmp_path / "test_sanitize.xml"
        path.write_bytes(xml_content.encode("utf-8"))

        text, meta = read_text(str(path))

        assert meta["sanitized"] is True
        # ANSI escape \x1b 제거됨
        assert "\x1b" not in text


# ---------------------------------------------------------------------------
# Step 2: LLM 게이트 테스트
# ---------------------------------------------------------------------------

class TestLLMGate:
    """preflight.run_llm_gate 의 동작.

    새 결정 행렬 (사용자 계약):
      LLM 가용:
        LLM OK  → 통과 (휴리스틱 참조 안 함)
        LLM NG  → PreflightError (휴리스틱 OK여도 LLM이 최종 결정권자)
      LLM 불가(예외 / client=None):
        휴리스틱 폴백 OK  → 통과
        휴리스틱 폴백 NG  → PreflightError
    """

    def test_ok_client_passes(self):
        """LLM OK → 통과 (휴리스틱 참조 안 함)."""
        from judge_tool.preflight import run_llm_gate

        meta = {}
        run_llm_gate("정상 텍스트 테스트 데이터", meta, OKClient(), "test.xml")
        assert meta.get("gate_decided_by") == "llm"

    def test_ng_client_raises_preflight_error_even_if_heuristic_ok(self):
        """LLM NG → PreflightError. 휴리스틱이 OK여도 LLM이 최종 결정권자.

        새 계약: LLM NG + 휴리스틱 OK 조합도 중단(기존 '오탐 무효화' 로직 제거).
        """
        from judge_tool.preflight import PreflightError, run_llm_gate

        # 정상 텍스트 → 휴리스틱은 OK이지만, LLM NG이므로 PreflightError
        normal_text = "정상적인 서버 점검 결과 텍스트 " * 20

        with pytest.raises(PreflightError) as exc_info:
            run_llm_gate(normal_text, {}, NGClient(), "test.xml")

        assert "NG" in str(exc_info.value)

    def test_ng_client_bad_text_raises_preflight_error(self):
        """LLM NG + 실제 손상 텍스트 → PreflightError."""
        from judge_tool.preflight import PreflightError, run_llm_gate

        bad_text = ("정상텍스트" * 10 + "■" * 20).replace("■", "�")

        with pytest.raises(PreflightError) as exc_info:
            run_llm_gate(bad_text, {}, NGClient(), "test.xml")

        assert "NG" in str(exc_info.value)

    def test_llm_exception_falls_back_to_heuristic_no_crash(self):
        """LLM 호출 예외 → 휴리스틱 폴백, 정상 텍스트는 크래시 없이 통과."""
        from judge_tool.preflight import run_llm_gate

        meta = {}
        run_llm_gate(
            "정상적인 서버 점검 결과 텍스트 " * 20,
            meta, ExceptionClient(), "test.xml")
        assert meta.get("gate_decided_by") == "heuristic(fallback)"

    def test_llm_exception_with_bad_text_heuristic_ng(self):
        """LLM 호출 예외 + 깨진 텍스트 → 휴리스틱 폴백 NG → PreflightError."""
        from judge_tool.preflight import PreflightError, run_llm_gate

        # replacement char 50% → 휴리스틱 NG
        bad_text = "가" * 10 + "�" * 10
        with pytest.raises(PreflightError):
            run_llm_gate(bad_text, {}, ExceptionClient(), "test.xml")

    def test_none_client_uses_heuristic(self):
        """client=None → 휴리스틱 폴백, 정상 텍스트 → 통과."""
        from judge_tool.preflight import run_llm_gate

        meta = {}
        run_llm_gate("정상 텍스트 내용 " * 20, meta, None, "test.xml")
        assert meta.get("gate_decided_by") == "heuristic(fallback)"

    def test_heuristic_rejects_short_text(self):
        """client=None + 너무 짧은 텍스트 → 휴리스틱 NG → PreflightError."""
        from judge_tool.preflight import PreflightError, run_llm_gate

        with pytest.raises(PreflightError):
            run_llm_gate("짧음", {}, None, "test.xml")

    def test_preflight_error_is_report_error(self):
        """PreflightError는 ReportError 하위 — main() 핸들러에서 포착됨."""
        from judge_tool.errors import ReportError
        from judge_tool.preflight import PreflightError

        assert issubclass(PreflightError, ReportError)

    def test_meta_records_gate_decided_by_llm(self):
        """LLM 성공 시 meta['gate_decided_by'] == 'llm'."""
        from judge_tool.preflight import run_llm_gate

        meta = {}
        run_llm_gate("정상 텍스트 테스트 데이터 " * 10, meta, OKClient(), "t.xml")
        assert meta["gate_decided_by"] == "llm"

    def test_meta_records_gate_decided_by_heuristic_on_exception(self):
        """LLM 예외 시 meta['gate_decided_by'] == 'heuristic(fallback)'."""
        from judge_tool.preflight import run_llm_gate

        meta = {}
        run_llm_gate("정상 텍스트 내용 " * 20, meta, ExceptionClient(), "t.xml")
        assert meta["gate_decided_by"] == "heuristic(fallback)"


# ---------------------------------------------------------------------------
# Step 3: run_preflight 통합 테스트
# ---------------------------------------------------------------------------

class TestRunPreflight:
    """preflight.run_preflight 의 통합 동작."""

    def test_skip_preflight_bypasses_gate(self, tmp_path):
        """skip=True → LLM 게이트 건너뜀, NG 클라이언트여도 통과."""
        from judge_tool.preflight import run_preflight

        xml_content = (
            '<?xml version="1.0" encoding="UTF-8"?>\n'
            "<script><results/></script>"
        )
        path = tmp_path / "test_skip.xml"
        path.write_bytes(xml_content.encode("utf-8"))

        # NG 클라이언트여도 skip=True이면 예외 없음
        text, meta = run_preflight(str(path), NGClient(), skip=True)
        assert isinstance(text, str)

    def test_normal_file_ok_client_passes(self, tmp_path):
        """정상 파일 + OK 클라이언트 → 예외 없이 교정 텍스트 반환."""
        from judge_tool.preflight import run_preflight

        xml_content = (
            '<?xml version="1.0" encoding="UTF-8"?>\n'
            "<script><results/></script>"
        )
        path = tmp_path / "test_normal.xml"
        path.write_bytes(xml_content.encode("utf-8"))

        text, meta = run_preflight(str(path), OKClient(), skip=False)
        assert "<script>" in text


# ---------------------------------------------------------------------------
# Step 4: main.run 통합 스모크 (--skip-preflight 동작, 기존 동작 불변)
# ---------------------------------------------------------------------------

class TestMainRunPreflight:
    """main.run 의 skip_preflight 플래그 동작 확인."""

    def _make_criteria(self, tmp_path):
        import openpyxl
        from judge_tool.profile import CLOUD

        criteria = os.path.join(str(tmp_path), "criteria.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = CLOUD.sheet_name
        rows = [
            ("PISM-001", "통신구간 암호화", 5, "스크립트", "방법1", "양호 기준"),
        ]
        for i, (iid, name, risk, etype, method, standard) in enumerate(rows):
            r = CLOUD.data_start_row + i
            ws.cell(r, CLOUD.id_col, iid)
            ws.cell(r, CLOUD.name_col, name)
            ws.cell(r, CLOUD.risk_col, risk)
            for vname in ("AWS", "Azure"):
                v = CLOUD.variants[vname]
                ws.cell(r, v.eval_type_col, etype)
                ws.cell(r, v.method_col, method)
                ws.cell(r, v.standard_col, standard)
        wb.save(criteria)
        return criteria

    def test_skip_preflight_flag_on_main_run(self, tmp_path):
        """skip_preflight=True로 main.run 호출 → 동작 불변, 예외 없음."""
        from judge_tool.main import run

        fixtures_dir = os.path.join(os.path.dirname(__file__), "fixtures")
        fixture_xml = os.path.join(fixtures_dir, "sample_aws_report.xml")

        report = os.path.join(str(tmp_path), "aws_report_test.xml")
        with open(fixture_xml, encoding="utf-8") as src:
            content = src.read()
        with open(report, "w", encoding="utf-8") as fh:
            fh.write(content)

        criteria = self._make_criteria(tmp_path)
        json_out = os.path.join(str(tmp_path), "result.json")
        xlsx_out = os.path.join(str(tmp_path), "result.xlsx")

        cov = run(
            report_path=report,
            criteria_path=criteria,
            profile_key="cloud",
            client=JudgeStubClient(),
            json_out=json_out,
            xlsx_out=xlsx_out,
            model_name="stub",
            skip_preflight=True,
        )
        assert os.path.exists(json_out)
        assert cov["judged"] >= 1

    def test_main_run_with_ok_gate_client(self, tmp_path):
        """OK 게이트 클라이언트로 main.run → 동작 불변."""
        from judge_tool.main import run

        fixtures_dir = os.path.join(os.path.dirname(__file__), "fixtures")
        fixture_xml = os.path.join(fixtures_dir, "sample_aws_report.xml")

        report = os.path.join(str(tmp_path), "aws_report_gate.xml")
        with open(fixture_xml, encoding="utf-8") as src:
            content = src.read()
        with open(report, "w", encoding="utf-8") as fh:
            fh.write(content)

        criteria = self._make_criteria(tmp_path)
        json_out = os.path.join(str(tmp_path), "result.json")
        xlsx_out = os.path.join(str(tmp_path), "result.xlsx")

        cov = run(
            report_path=report,
            criteria_path=criteria,
            profile_key="cloud",
            client=JudgeStubClient(),
            json_out=json_out,
            xlsx_out=xlsx_out,
            model_name="stub",
            skip_preflight=False,
        )
        assert os.path.exists(json_out)
        assert cov["judged"] >= 1

    def test_main_run_ng_gate_raises_preflight_error(self, tmp_path):
        """NG 게이트 + skip_preflight=False → PreflightError.

        새 계약: client=None → 휴리스틱 폴백만 사용.
        replacement char 과다 파일 → 휴리스틱 NG → PreflightError.
        """
        from judge_tool.errors import ReportError
        from judge_tool.preflight import PreflightError

        # replacement char 과다 파일 — 휴리스틱 NG 확정 (LLM 필요 없음)
        xml_content = (
            '<?xml version="1.0" encoding="UTF-8"?>\n'
            "<script><results><dump><items><id>aws_report_synth</id></items>"
            "<output><![CDATA["
            # 30% replacement char → 5% 임계 초과 → heuristic NG
            + ("정상 " * 10 + "■" * 15) * 5
            + "]]></output></dump></results></script>"
        )
        report = os.path.join(str(tmp_path), "aws_report_ng.xml")
        # replacement char �를 직접 삽입
        broken_content = xml_content.replace("■", "�")
        with open(report, "w", encoding="utf-8") as fh:
            fh.write(broken_content)

        criteria = self._make_criteria(tmp_path)
        json_out = os.path.join(str(tmp_path), "result.json")
        xlsx_out = os.path.join(str(tmp_path), "result.xlsx")

        # client=None → 휴리스틱만 사용 → replacement char 과다 → NG
        with pytest.raises((PreflightError, ReportError)):
            from judge_tool.main import run
            run(
                report_path=report,
                criteria_path=criteria,
                profile_key="cloud",
                client=None,
                json_out=json_out,
                xlsx_out=xlsx_out,
                model_name="stub",
                skip_preflight=False,
            )


# ---------------------------------------------------------------------------
# Step 5: parser 통합 — 컨테이너 XML mojibake before/after
# ---------------------------------------------------------------------------

class TestParserEncodingCorrection:
    """파서(_read_text 경유)의 인코딩 교정 실제 동작 확인."""

    def test_container_xml_parse_korean_correct(self):
        """container_xml.parse()가 EUC-KR 선언+UTF-8 실제 파일을 올바르게 처리.

        교정 전 (buggy): euc-kr decode → mojibake → LLM에 깨진 텍스트 입력.
        교정 후 (fixed): utf-8 strict detect → 한글 정상 → LLM에 정상 텍스트.
        """
        if not os.path.exists(CONTAINER_XML):
            pytest.skip("컨테이너 실데이터 파일이 없어 스킵")

        from judge_tool.parsers.container_xml import parse

        results = parse(CONTAINER_XML)
        assert len(results) > 0, "최소 1개 이상의 결과가 있어야 함"

        # 모든 evidence를 합친 뒤 한글 포함 여부 확인
        all_evidence = " ".join(
            ev.evidence for (cid, evs, _) in results for ev in evs
        )

        # 한글이 정상 (mojibake 마커 없음)
        replacement_count = all_evidence.count("�")
        assert replacement_count == 0, \
            f"evidence에 replacement char가 없어야 함, found: {replacement_count}"

        # 실제 한글 텍스트가 포함되어 있음
        korean_words = re.findall(r"[가-힣]{2,}", all_evidence)
        assert len(korean_words) > 0, \
            "교정 후 evidence에 한글 단어가 포함되어야 함"

    def test_server_xml_normal_file_unchanged(self):
        """정상 인코딩 서버 파일 → 기존 동작 불변."""
        linux_xml = os.path.join(
            PROJECT_ROOT, "collected", "server", "linux", "linux-s-sample.xml")
        if not os.path.exists(linux_xml):
            pytest.skip("서버 실데이터 파일이 없어 스킵")

        from judge_tool.parsers.server_xml import parse

        results = parse(linux_xml)
        assert len(results) > 0, "정상 파일은 parse()가 결과를 반환해야 함"

        # evidence가 손상되지 않았는지 확인
        all_evidence = " ".join(
            ev.evidence for (cid, evs, _) in results for ev in evs
        )
        replacement_count = all_evidence.count("�")
        assert replacement_count == 0, \
            "정상 인코딩 파일에는 replacement char가 없어야 함"
