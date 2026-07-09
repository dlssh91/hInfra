import openpyxl
from judge_tool.criteria_loader import load_criteria
from judge_tool.profile import DB_MYSQL, DB_ORACLE, DB_MSSQL, DB_POSTGRESQL


def _make_db_xlsx(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "데이터베이스"
    ws.cell(4, 2, "평가항목ID"); ws.cell(4, 7, "평가항목"); ws.cell(4, 8, "위험도")
    ws.cell(4, 17, "평가대상(AWS RDS MYSQL)")
    ws.cell(4, 37, "판단기준(RDS)"); ws.cell(4, 38, "판단방법(RDS)")
    # row5: rds 적용(o) + 판단기준 있음 → judgeable
    ws.cell(5, 2, "DBM-001"); ws.cell(5, 7, "비밀번호"); ws.cell(5, 8, 4.0)
    ws.cell(5, 17, "o"); ws.cell(5, 37, "* 양호 - ...\n* 취약 - ..."); ws.cell(5, 38, "확인방법")
    # row6: rds 적용 아님(공란) → not applicable
    ws.cell(6, 2, "DBM-099"); ws.cell(6, 7, "해당없음"); ws.cell(6, 8, 3.0)
    ws.cell(6, 37, "* 양호 - ..."); ws.cell(6, 38, "방법")
    wb.save(path)


def test_db_applicable_from_marker(tmp_path):
    p = str(tmp_path / "db.xlsx")
    _make_db_xlsx(p)
    crit = load_criteria(p, DB_MYSQL)
    c1 = crit[("DBM-001", "mysql_rds")]
    assert c1.applicable is True and c1.is_judgeable is True
    assert "양호" in c1.standard
    c99 = crit[("DBM-099", "mysql_rds")]
    assert c99.applicable is False and c99.is_judgeable is False


# ── 네이티브 변형 라벨 정합성(실제 기준 xlsx 사용) ──────────────────────────

def test_native_dbm001_label_c(criteria_xlsx_path):
    """네이티브에서 새로 평가대상이 되는 DBM-001은 mssql/pg 모두 label C여야
    한다(default A로 LLM 오판정되지 않도록 yaml에 추가)."""
    mssql = load_criteria(criteria_xlsx_path, DB_MSSQL, profile_key="db_mssql")
    assert mssql[("DBM-001", "mssql_native")].label == "C"
    pg = load_criteria(criteria_xlsx_path, DB_POSTGRESQL, profile_key="db_postgresql")
    assert pg[("DBM-001", "pg_native")].label == "C"


def test_pg_native_dbm007_override_a(criteria_xlsx_path):
    """PG DBM-007은 클라우드 전용 C(passwordcheck 설치불가)이나, pg_native는
    variant 오버라이드로 A(LLM 판정)여야 한다."""
    pg = load_criteria(criteria_xlsx_path, DB_POSTGRESQL, profile_key="db_postgresql")
    assert pg[("DBM-007", "pg_native")].label == "A"
    # 클라우드 변형은 여전히 C(오버라이드 영향 없음).
    assert pg[("DBM-007", "pg_rds")].label == "C"


# ── F2. DBM-025(EoS) det_common 자동판정 제거 — eol.yaml 권위경로 회귀 ───────

def test_native_dbm025_classify_as_det_not_det_common(criteria_xlsx_path):
    """DBM-025(EoS DB 사용)는 label D이므로 item_configs가 judgment_method를
    지정하지 않으면 classify_method가 'det'를 도출해야 한다(_defer_or_eol →
    judge_eol의 eol.yaml as_of/staleness 권위경로). 노후 지식 기반
    'det_common' 자동판정(예: oracle 룰이 '12'만 EoS로 알아 12.1/12.2/18c를
    거짓양호 처리)으로 새지 않는지 4개 native 엔진 전부에서 고정한다.
    mariadb는 이미 이 형태(judgment_method 미지정)이다."""
    cases = [
        (DB_MYSQL, "db_mysql", "mysql_native"),
        (DB_ORACLE, "db_oracle", "oracle_native"),
        (DB_MSSQL, "db_mssql", "mssql_native"),
        (DB_POSTGRESQL, "db_postgresql", "pg_native"),
    ]
    for profile, key, variant in cases:
        crit = load_criteria(criteria_xlsx_path, profile, profile_key=key)
        c = crit[("DBM-025", variant)]
        assert c.judgment_method == "det", (
            f"{key}/{variant} DBM-025: judgment_method={c.judgment_method!r} "
            f"(기대: 'det' — det_common 노후지식 자동판정이 아닌 eol.yaml 권위경로)"
        )
        assert c.judgment_method != "det_common"
        assert c.label == "D"
