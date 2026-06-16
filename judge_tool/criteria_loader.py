import os
from typing import Dict, Tuple

import openpyxl
import yaml

from judge_tool.models import Criterion
from judge_tool.profile import Profile

_ITEM_CONFIG_DIR = os.path.join(os.path.dirname(__file__), "item_configs")


def _load_item_configs(profile_key: str) -> Dict:
    """item_configs/{profile_key}.yaml 로딩. 파일 없으면 빈 dict 반환."""
    path = os.path.join(_ITEM_CONFIG_DIR, f"{profile_key}.yaml")
    if not os.path.exists(path):
        return {}
    with open(path, encoding="utf-8") as fh:
        data = yaml.safe_load(fh) or {}
    return data


def classify_method(label: str, *, has_summary: bool,
                    in_empty_means_good: bool) -> str:
    """라벨+설정에서 판단방식 5분류(정적)를 도출하는 단일 출처.

    - C/D            → det            (canned 기술한계 / EOL·패치 결정론)
    - B + 요약지시    → interview       (인터뷰-내용정리)
    - B + 요약지시 X  → interview_holdonly (인터뷰-내용정리X, 요약 없이 보류만)
    - A + empty_means_good → llm_det   (빈결과=양호 결정론 가드 + LLM)
    - A (그 외)       → llm

    하이브리드(B + empty_means_good)는 주(主)방식 기준 interview로 분류한다.
    empty_means_good 가드는 빈결과 시 verdict/rationale에 직교적으로 드러난다.
    """
    if label in ("C", "D"):
        return "det"
    if label == "B":
        return "interview" if has_summary else "interview_holdonly"
    # label A (및 미지 라벨)은 LLM 판정. empty_means_good이면 결정론 가드 동반.
    return "llm_det" if in_empty_means_good else "llm"


def _cell(ws, row, col) -> str:
    v = ws.cell(row=row, column=col).value
    return "" if v is None else str(v).strip()


def load_criteria(xlsx_path: str,
                  profile: Profile,
                  profile_key: str = "") -> Dict[Tuple[str, str], Criterion]:
    """xlsx → {(item_id, variant): Criterion}. 모든 변형을 로드한다."""
    item_configs = _load_item_configs(profile_key) if profile_key else {}
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    try:
        ws = wb[profile.sheet_name]
        out: Dict[Tuple[str, str], Criterion] = {}

        for row in range(profile.data_start_row, ws.max_row + 1):
            item_id = _cell(ws, row, profile.id_col)
            if not item_id:
                continue
            # mapper.aggregate 가 증거 측 CheckID 를 normalize_id 로 정규화하므로
            # 기준 키도 동일 정규형으로 맞춰 join 비대칭을 제거한다.
            item_id = profile.normalize_id(item_id)
            name = _cell(ws, row, profile.name_col)
            # risk는 숫자 원본이 필요하므로 _cell(str 변환) 대신 .value를 직접 사용
            risk_raw = ws.cell(row=row, column=profile.risk_col).value
            try:
                risk = float(risk_raw)
            except (TypeError, ValueError):
                risk = None

            for vname, vspec in profile.variants.items():
                eval_type = (_cell(ws, row, vspec.eval_type_col)
                             if vspec.eval_type_col is not None else "")
                standard = _cell(ws, row, vspec.standard_col)
                method = _cell(ws, row, vspec.method_col)
                if vspec.applicability_col is not None:
                    # DB: 평가대상 컬럼 'o'
                    marker = _cell(ws, row, vspec.applicability_col).strip().lower()
                    applicable = (marker == "o")
                elif vspec.applies_when_standard:
                    # network generic: 벤더중립 판단기준(C18)이 있으면 적용
                    applicable = bool(standard)
                else:
                    # cloud: 기존 is_judgeable 동치(스크립트 기반 & N/A 아님)
                    applicable = ("스크립트" in eval_type) and eval_type != "N/A"
                cfg = item_configs.get(item_id, {})
                # variant별 오버라이드: cfg["variants"][vname]이 있으면 해당 값 우선
                vcfg = cfg.get("variants", {}).get(vname, {})
                label = vcfg.get("label", cfg.get("label", "A"))
                summary_instruction = vcfg.get(
                    "summary_instruction", cfg.get("summary_instruction"))
                # yaml 직접 지정(예: fw_policy)이 classify_method 결과보다 우선.
                yaml_method = vcfg.get("judgment_method",
                                       cfg.get("judgment_method"))
                if yaml_method:
                    judgment_method = yaml_method
                else:
                    judgment_method = classify_method(
                        label,
                        has_summary=bool(summary_instruction),
                        in_empty_means_good=(item_id in profile.empty_means_good))
                # §6 임계값 로딩: yaml thresholds/thresholds_source → Criterion.
                # 없으면 {} / None(기존 동작 불변 — 아직 어떤 yaml도 thresholds 미보유).
                thresholds_raw = vcfg.get("thresholds", cfg.get("thresholds"))
                thresholds = dict(thresholds_raw) if isinstance(thresholds_raw, dict) else {}
                thresholds_source = vcfg.get("thresholds_source",
                                             cfg.get("thresholds_source"))
                out[(item_id, vname)] = Criterion(
                    item_id=item_id,
                    item_name=name,
                    risk=risk,
                    variant=vname,
                    eval_type=eval_type,
                    standard=standard,
                    method=method,
                    applicable=applicable,
                    label=label,
                    canned_message=vcfg.get("canned_message",
                                            cfg.get("canned_message")),
                    summary_instruction=summary_instruction,
                    eol_check=bool(vcfg.get("eol_check",
                                            cfg.get("eol_check", False))),
                    patch_check=bool(vcfg.get("patch_check",
                                              cfg.get("patch_check", False))),
                    judgment_method=judgment_method,
                    thresholds=thresholds,
                    thresholds_source=thresholds_source,
                )
        return out
    finally:
        wb.close()
