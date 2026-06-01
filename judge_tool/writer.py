import json
from dataclasses import asdict
from typing import Dict, List, Tuple

import openpyxl
from openpyxl.styles import Font, PatternFill

from judge_tool.models import Criterion, Judgment


def build_coverage(criteria: Dict[Tuple[str, str], Criterion],
                   judgments: List[Judgment], variant: str) -> Dict:
    """스크립트 기반·해당 variant 항목 중 미판정 목록 산출."""
    expected_ids = {
        c.item_id for (iid, v), c in criteria.items()
        if v == variant and c.is_script_based and c.eval_type != "N/A"
    }
    judged_ids = {j.item_id for j in judgments}
    missing = sorted(expected_ids - judged_ids)
    return {"expected": len(expected_ids), "judged": len(judged_ids),
            "missing": missing}


def write_json(judgments: List[Judgment], meta: Dict, coverage: Dict,
               path: str) -> None:
    payload = {
        "metadata": meta,
        "coverage": coverage,
        "judgments": [asdict(j) for j in judgments],
    }
    with open(path, "w", encoding="utf-8") as fh:
        json.dump(payload, fh, ensure_ascii=False, indent=2)


_HEADERS = ["항목ID", "항목명", "위험도", "변형", "판정", "확신도",
            "스크립트status", "일치여부", "재검토", "범위", "관리체계검토",
            "근거", "인용증거"]
_REVIEW_FILL = PatternFill("solid", fgColor="FFF2CC")  # 연노랑


def write_excel(judgments: List[Judgment], meta: Dict, coverage: Dict,
                path: str) -> None:
    wb = openpyxl.Workbook()

    # 요약 시트
    summary = wb.active
    summary.title = "요약"
    rows = [
        ("도구 버전", meta.get("tool_version", "")),
        ("기준 버전", meta.get("criteria_version", "")),
        ("사용 모델", meta.get("model", "")),
        ("생성 시각", meta.get("generated_at", "")),
        ("입력 파일", meta.get("source_file", "")),
        ("입력 SHA256", meta.get("source_sha256", "")),
        ("대상 항목수", coverage.get("expected", "")),
        ("판정 항목수", coverage.get("judged", "")),
        ("미판정 항목", ", ".join(coverage.get("missing", []))),
        ("재검토 필요수", sum(1 for j in judgments if j.needs_review)),
    ]
    for r, (k, v) in enumerate(rows, start=1):
        summary.cell(r, 1, k).font = Font(bold=True)
        summary.cell(r, 2, v)

    # 판정결과 시트
    ws = wb.create_sheet("판정결과")
    ws.append(_HEADERS)
    for c in ws[1]:
        c.font = Font(bold=True)
    for j in judgments:
        ws.append([
            j.item_id, j.item_name, j.risk, j.variant, j.verdict,
            round(j.confidence, 2), j.script_status, j.agreement,
            "예" if j.needs_review else "", j.scope,
            "예" if j.management_review_needed else "",
            j.rationale, " | ".join(j.cited_evidence),
        ])
        if j.needs_review:
            for c in ws[ws.max_row]:
                c.fill = _REVIEW_FILL
    wb.save(path)
